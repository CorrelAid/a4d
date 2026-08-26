"""Measure the taxonomy's suspected blind spots against real trackers.

Ticket 70 asks what can go wrong in a workbook that the pipeline never
reports. :mod:`scripts.finding_inventory` finds the *candidates* in the source
-- places that discard or replace a value in a function that emits no finding.
This script sizes them on real data, because a blind spot nobody can put a
number on is a code review, not a finding.

Each probe is named for the gap it measures and prints its own evidence.

Usage::

    uv run python scripts/finding_blind_spots.py --data-root "/Volumes/.../a4dphase2_upload"
    uv run python scripts/finding_blind_spots.py --probe sheets
"""

from __future__ import annotations

import argparse
import calendar
from collections import Counter
from pathlib import Path

import polars as pl
from openpyxl import load_workbook

MONTH_ABBRS = tuple(calendar.month_abbr)[1:]
STATIC_SHEETS = ("Patient List", "Annual")


def tracker_files(data_root: Path) -> list[Path]:
    """Every workbook the pipeline would treat as a tracker.

    Mirrors the discovery rule rather than importing it: this script must be
    able to say "the pipeline never looked at this sheet", which means walking
    the same files without going through the extractor.
    """
    return sorted(
        p
        for p in data_root.rglob("*.xlsx")
        if not p.name.startswith("~$") and "output" not in p.relative_to(data_root).parts
    )


def probe_sheets(data_root: Path) -> None:
    """Sheets no code path ever opens.

    ``find_month_sheets`` keeps only sheets whose name *starts with* a
    capitalised month abbreviation, and the static sheets are matched by exact
    string. Anything else is skipped before any finding could be emitted --
    including, by construction, an all-caps or misspelled month name.
    """
    unseen: Counter[str] = Counter()
    unseen_with_data: list[tuple[str, str, int]] = []
    files = tracker_files(data_root)
    for path in files:
        try:
            wb = load_workbook(
                path, read_only=True, data_only=True, keep_vba=False, keep_links=False
            )
        except Exception as exc:  # a workbook the pipeline also could not open
            print(f"  unreadable: {path.name}: {str(exc)[:70]}")
            continue
        for name in wb.sheetnames:
            if any(name.startswith(abbr) for abbr in MONTH_ABBRS) or name in STATIC_SHEETS:
                continue
            unseen[name] += 1
            rows = wb[name].max_row or 0
            if rows > 1:
                unseen_with_data.append((path.name, name, rows))
        wb.close()

    print(f"{len(files)} workbooks scanned")
    print(f"{sum(unseen.values())} sheets are never opened, {len(unseen)} distinct names")
    print(f"{len(unseen_with_data)} of them hold more than a header row\n")
    for name, count in unseen.most_common(40):
        print(f"  {count:>4}x  {name!r}")
    print("\n  looks like a month name the matcher missed:")
    for name, count in unseen.most_common():
        lowered = name.lower()
        if any(lowered.startswith(abbr.lower()) for abbr in MONTH_ABBRS):
            print(f"  {count:>4}x  {name!r}  <-- case-sensitive match dropped this")


def probe_sex(data_root: Path) -> None:
    """Values ``fix_sex`` replaced with the ``Undefined`` sentinel, silently."""
    cleaned = data_root / "output" / "patient_data_cleaned"
    frame = pl.scan_parquet(cleaned / "*.parquet").select("file_name", "sex").collect()
    print(frame["sex"].value_counts(sort=True))


def probe_t1d_age(data_root: Path) -> None:
    """Rows where the diagnosis date precedes the date of birth.

    ``_fix_age_from_dob`` reports the same contradiction on the visit age as
    ``age_negative_from_dob``; ``_fix_t1d_diagnosis_age`` nulls it and says
    nothing.
    """
    cleaned = data_root / "output" / "patient_data_cleaned"
    frame = (
        pl.scan_parquet(cleaned / "*.parquet")
        .select("file_name", "patient_id", "dob", "t1d_diagnosis_date", "t1d_diagnosis_age")
        .filter(
            pl.col("dob").is_not_null()
            & pl.col("t1d_diagnosis_date").is_not_null()
            & (pl.col("t1d_diagnosis_date") < pl.col("dob"))
        )
        .collect()
    )
    print(f"{frame.height} rows, {frame['file_name'].n_unique()} trackers")
    print(frame.head(10))


def probe_product_recipient(data_root: Path) -> None:
    """Stock released to a patient ID no patient sheet in that tracker knows.

    ``link_product_patient`` counts these and writes them to DEBUG logs only,
    so they reach neither the findings table nor the report.
    """
    tables = data_root / "output" / "tables"
    product = pl.read_parquet(tables / "product_data.parquet")
    patients = (
        pl.read_parquet(tables / "patient_data_monthly.parquet")
        .select("file_name", "patient_id")
        .unique()
        .with_columns(pl.lit(True).alias("_matched"))
    )
    candidates = product.filter(
        pl.col("product_released_to").is_not_null() & (pl.col("product_released_to") != "Undefined")
    )
    unmatched = candidates.join(
        patients,
        left_on=["file_name", "product_released_to"],
        right_on=["file_name", "patient_id"],
        how="left",
    ).filter(pl.col("_matched").is_null())
    print(f"{candidates.height} rows name a recipient")
    print(
        f"{unmatched.height} name one no patient sheet in that tracker knows, "
        f"across {unmatched['file_name'].n_unique()} trackers "
        f"and {unmatched['product_released_to'].n_unique()} distinct IDs"
    )
    print(
        unmatched.group_by("file_name", "product_released_to")
        .agg(pl.len().alias("rows"))
        .sort("rows", descending=True)
        .head(15)
    )


def probe_static_dropped_rows(data_root: Path) -> None:
    """Patient List / Annual rows dropped for a null or ``#``-error ID.

    The month-sheet path reports the same two defects as
    ``missing_required_field`` and ``excel_error_patient_id``; the static-sheet
    path filters them out without a word.
    """
    null_ids = 0
    error_ids = 0
    affected: set[str] = set()
    for path in tracker_files(data_root):
        try:
            wb = load_workbook(
                path, read_only=True, data_only=True, keep_vba=False, keep_links=False
            )
        except Exception:
            continue
        for sheet in STATIC_SHEETS:
            if sheet not in wb.sheetnames:
                continue
            rows = list(wb[sheet].iter_rows(values_only=True))
            if not rows:
                continue
            # The header is not row 1: these sheets open with a clinic banner
            # and a merged section title, so the real header sits several rows
            # down (row 9 in the 2026 template). Find it the way a reader
            # would -- the first row carrying a "Patient ID" cell.
            header_row = None
            id_col = None
            for index, row in enumerate(rows[:30]):
                for col, cell in enumerate(row):
                    text = str(cell).strip().lower().replace("\n", " ") if cell else ""
                    if "patient" in text and "id" in text:
                        header_row, id_col = index, col
                        break
                if header_row is not None:
                    break
            if id_col is None or header_row is None:
                continue
            for row in rows[header_row + 1 :]:
                if id_col >= len(row):
                    continue
                value = row[id_col]
                if value is None or str(value).strip() == "":
                    if any(c is not None and str(c).strip() for c in row):
                        null_ids += 1
                        affected.add(path.name)
                elif str(value).startswith("#"):
                    error_ids += 1
                    affected.add(path.name)
        wb.close()
    print(f"{null_ids} static-sheet rows carry data but no patient ID")
    print(f"{error_ids} carry an Excel error where the ID belongs")
    print(f"across {len(affected)} trackers -- none of them reaches a finding")


PROBES = {
    "sheets": probe_sheets,
    "sex": probe_sex,
    "t1d-age": probe_t1d_age,
    "product-recipient": probe_product_recipient,
    "static-rows": probe_static_dropped_rows,
}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data-root",
        type=Path,
        default=Path("/Volumes/USB SanDisk 3.2Gen1 Media/a4d/a4dphase2_upload"),
    )
    parser.add_argument("--probe", choices=sorted(PROBES), action="append")
    args = parser.parse_args()

    if not args.data_root.exists():
        parser.error(f"data root not found: {args.data_root}")

    for name in args.probe or sorted(PROBES):
        print(f"\n=== {name} ===")
        PROBES[name](args.data_root)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
