"""Unit tests for patient extraction helper functions."""

import datetime
import random
from unittest.mock import Mock

import polars as pl
import pytest
from openpyxl import Workbook

from a4d.extract.patient import (
    filter_valid_columns,
    find_data_start_row,
    find_dropped_data_columns,
    find_layout_changes,
    join_static_sheet,
    merge_headers,
    read_header_rows,
    read_patient_rows,
    recover_blank_headers,
    report_duplicate_patients_in_sheet,
)
from a4d.findings import findings_collected


def create_mock_mapper(known_columns: set[str], standard_names: dict[str, str] | None = None):
    """Create a mock ColumnMapper that validates specific column names."""
    mapper = Mock()
    mapper.is_known_column = lambda col: col in known_columns
    names = standard_names or {}
    mapper.get_standard_name = lambda col: names.get(col, col)
    return mapper


class TestFindDataStartRow:
    """Tests for find_data_start_row() function."""

    def test_data_starts_at_row_1(self):
        """Test when data starts at the very first row."""
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 1
        ws["A2"] = 2

        result = find_data_start_row(ws)
        assert result == 1

        wb.close()

    def test_data_starts_after_empty_rows(self):
        """Test when there are empty rows before data."""
        wb = Workbook()
        ws = wb.active
        # Leave rows 1-10 empty
        ws["A11"] = 1
        ws["A12"] = 2

        result = find_data_start_row(ws)
        assert result == 11

        wb.close()

    def test_realistic_tracker_layout(self):
        """Test with realistic tracker layout (headers at rows 75-76, data at 77)."""
        wb = Workbook()
        ws = wb.active

        # Simulate typical tracker: empty rows, then title rows, then headers, then data
        # Title area NOT in column A (column A stays empty until headers)
        ws["B1"] = "Hospital Name"
        ws["C1"] = "General Hospital"

        # Headers at rows 75-76 (typical for real trackers)
        ws["B75"] = "Patient"
        ws["B76"] = "ID*"

        # Data starts at row 77
        ws["A77"] = 1
        ws["A78"] = 2

        result = find_data_start_row(ws)
        assert result == 77  # First non-None in column A

        wb.close()

    def test_randomized_data_position(self):
        """Test with randomized data start position."""
        wb = Workbook()
        ws = wb.active

        # Random start position between 10 and 100
        random_start = random.randint(10, 100)

        # Insert first data value at random position (must be numeric)
        ws[f"A{random_start}"] = 1

        result = find_data_start_row(ws)
        assert result == random_start

        wb.close()

    def test_column_a_empty_raises_error(self):
        """Test that ValueError is raised when column A is empty."""
        wb = Workbook()
        ws = wb.active

        # Put data in other columns but not A
        ws["B1"] = "Some data"
        ws["C5"] = "More data"

        with pytest.raises(ValueError, match="No patient data found in column A"):
            find_data_start_row(ws)

        wb.close()

    def test_ignores_none_values(self):
        """Test that None/empty cells are skipped correctly."""
        wb = Workbook()
        ws = wb.active

        # Explicitly set some cells to None (they start as None anyway)
        ws["A1"] = None
        ws["A2"] = None
        ws["A3"] = None
        ws["A4"] = 1  # First numeric data

        result = find_data_start_row(ws)
        assert result == 4

        wb.close()

    def test_blank_string_row_number_directly_above_data_is_data(self):
        """A cleared row-number cell must not push the start past its own data row.

        2022 Children's Hospital 2, Oct22: the first patient's row number was
        deleted, leaving a whitespace-only string. Starting one row later reads
        that patient's row as a header row, which costs the whole sheet.
        """
        wb = Workbook()
        ws = wb.active

        ws["B68"] = "Patient ID*"
        ws["A70"] = " "  # row number cleared, patient data still present
        ws["B70"] = "VN_CH001"
        ws["A71"] = 2
        ws["B71"] = "VN_CH002"

        result = find_data_start_row(ws)
        assert result == 70

        wb.close()

    def test_blank_string_far_above_data_is_not_data(self):
        """A whitespace-only cell separated from the data block stays skipped.

        NOGH 2026 and Phattalung 2021 carry one ~20 rows above the patients;
        it is layout residue, not a row number, and the header rows sit between.
        """
        wb = Workbook()
        ws = wb.active

        ws["A29"] = ""  # residue well above the data block
        ws["A48"] = 1
        ws["A49"] = 2

        result = find_data_start_row(ws)
        assert result == 48

        wb.close()

    def test_leading_text_in_column_a_is_not_data(self):
        """A stray word in column A row 1 must not be read as the data start.

        2025/2026 VNCH and 2026 Gensan hold a bare 'm'/'f'/'n' in A1, which is
        why the start cannot simply be the first non-empty cell in column A --
        measured across all 254 trackers, that rule would start 14 sheets here.
        """
        wb = Workbook()
        ws = wb.active

        ws["A1"] = "m"
        ws["A84"] = 1

        result = find_data_start_row(ws)
        assert result == 84

        wb.close()

    def test_scans_read_only_worksheet_in_one_pass(self, tmp_path):
        """Must not re-parse the sheet's XML on every row (O(n^2) on read-only sheets).

        ReadOnlyWorksheet.cell() re-parses from row 1 on every call, so a
        per-row .cell() loop is quadratic in the row count before data starts.
        Asserts _cells_by_row (the underlying XML scan) runs at most once,
        which only a single iter_rows() pass over column A can achieve.
        """
        from unittest.mock import patch

        from openpyxl import Workbook, load_workbook
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet

        path = tmp_path / "dense.xlsx"
        wb = Workbook()
        ws = wb.active
        for row in range(1, 200):
            for col in range(1, 40):
                ws.cell(row, col, value=f"text{row}_{col}")
        ws["A200"] = 1
        wb.save(path)
        wb.close()

        wb2 = load_workbook(path, read_only=True)
        ws2 = wb2.active

        with patch.object(ReadOnlyWorksheet, "_cells_by_row", wraps=ws2._cells_by_row) as spy:
            result = find_data_start_row(ws2)

        assert result == 200
        assert spy.call_count == 1

        wb2.close()


class TestReadPatientRows:
    """Tests for read_patient_rows() row-acceptance rules."""

    @staticmethod
    def _sheet(rows):
        wb = Workbook()
        ws = wb.active
        for r, values in enumerate(rows, start=1):
            for c, value in enumerate(values, start=1):
                if value is not None:
                    ws.cell(row=r, column=c, value=value)
        return wb, ws

    def test_keeps_a_numbered_row(self):
        wb, ws = self._sheet([[1, "VN_VC001", "VN_VC001", "Y"]])

        assert len(read_patient_rows(ws, 1, 4)) == 1

        wb.close()

    def test_keeps_an_unnumbered_row_that_carries_data(self):
        """2024_Mahosot Jun24 row 380: a real patient record whose row number
        was never filled in. Bounding the data block by the row-number column
        alone would silently discard it, which is why an unnumbered row that
        carries actual data is kept."""
        wb, ws = self._sheet([[None, "LA-MH088", "LA-MH088", "Y", "2024-06-19"]])

        assert len(read_patient_rows(ws, 1, 5)) == 1

        wb.close()

    def test_drops_an_unnumbered_row_that_only_repeats_its_own_identifier(self):
        """2024_Vietnam National Children Jul24 rows 165-188: a bare list of
        patient IDs left below the data block, two non-empty cells per row and
        nothing else. Treating those as monthly records invented 24 patients."""
        wb, ws = self._sheet([[None, "VN_VC052", "VN_VC052", None, None]])

        assert read_patient_rows(ws, 1, 5) == []

        wb.close()

    def test_stops_at_the_first_completely_empty_row(self):
        wb, ws = self._sheet(
            [[1, "VN_VC001", "VN_VC001"], [None, None, None], [2, "VN_VC002", "VN_VC002"]]
        )

        assert len(read_patient_rows(ws, 1, 3)) == 1

        wb.close()

    def test_reports_data_left_below_the_blank_row(self, collector):
        """2026 Preah Kossamak, Patient List/May26/Jun26: the clinic left a gap,
        wrote a 'PENDING TRANSFER KBH' banner, and started a second numbered
        block below it. Reading stops at the gap, which is correct -- those 14
        patients are still Kantha Bopha's and are in that tracker under their
        own IDs -- but it used to happen in total silence."""
        wb, ws = self._sheet(
            [
                [1, "KH_KB001_PK", "KH_KB001_PK"],
                [None, None, None],
                [None, "PENDING TRANSFER KBH", None],
                [1, "KH_KB114_PK", "KH_KB114_PK"],
            ]
        )
        ws.title = "Jun26"

        assert len(read_patient_rows(ws, 1, 3)) == 1

        reported = collector.to_dataframe()
        row = reported.filter(pl.col("error_code") == "data_below_blank_row")
        assert len(row) == 1
        assert row["sheet_name"][0] == "Jun26"
        assert "1 row(s) of data" in row["message"][0]
        assert "row 2" in row["message"][0]

        wb.close()

    def test_stays_silent_when_nothing_readable_follows_the_blank_row(self):
        """A stray note below the block is not a patient record: without a row
        number or an identifier it would be skipped even if reading continued,
        so the blank row costs nothing and there is nothing to report."""
        wb, ws = self._sheet(
            [
                [1, "VN_VC001", "VN_VC001"],
                [None, None, None],
                [None, None, "checked by Dr. L"],
            ]
        )

        with findings_collected(file_name="t") as c:
            read_patient_rows(ws, 1, 3)

        assert c.to_dataframe().is_empty()

        wb.close()

    def test_skips_a_row_with_neither_number_nor_identifier(self):
        wb, ws = self._sheet([[None, None, "stray note"], [1, "VN_VC001", "VN_VC001"]])

        rows = read_patient_rows(ws, 1, 3)

        assert [r[1] for r in rows] == ["VN_VC001"]

        wb.close()

    def test_recovers_a_number_typed_into_a_date_formatted_cell(self):
        """2025 Hat Yai Annual!H, TH_HY035: the systolic cell carries a
        dd-mmm-yyyy format, so Excel stored the typed 120 as 29-Apr-1900 and
        openpyxl hands back that datetime. The datum is the number."""
        wb, ws = self._sheet([[1, "TH_HY035", "TH_HY035", None]])
        ws.cell(row=1, column=4, value=datetime.datetime(1900, 4, 29))

        rows = read_patient_rows(ws, 1, 4)

        assert rows[0][3] == 120

        wb.close()

    def test_recovers_a_bare_year_typed_into_a_date_formatted_cell(self):
        """2019 Yangon Children's Patient List!F, MM_YC005: the D.O.B. cell
        holds the bare year 2003 under a date format, so openpyxl resolves
        serial 2003 to 25-Jun-1905 and the year is lost before cleaning can
        read it (ticket 52). Its 2017, 2018 and 2020 siblings leave the same
        cell General-formatted, so only this one file was affected."""
        wb, ws = self._sheet([[1, "MM_YC005", "MM_YC005", None]])
        ws.cell(row=1, column=4, value=datetime.datetime(1905, 6, 25))

        rows = read_patient_rows(ws, 1, 4)

        assert rows[0][3] == 2003

        wb.close()

    def test_keeps_a_date_that_could_plausibly_have_been_typed(self):
        """The inverse case, already settled for product as
        openpyxl_date_typed_stray_cell: a real date landing in the wrong
        column stays a date."""
        wb, ws = self._sheet([[1, "TH_CP005", "TH_CP005", None]])
        ws.cell(row=1, column=4, value=datetime.datetime(1956, 8, 1))

        rows = read_patient_rows(ws, 1, 4)

        assert rows[0][3] == datetime.datetime(1956, 8, 1)

        wb.close()


class TestReadHeaderRows:
    """Tests for read_header_rows() function."""

    def test_basic_two_row_headers(self):
        """Test reading basic two-row headers."""
        wb = Workbook()
        ws = wb.active

        # Data starts at row 5, so headers are at rows 3 and 4
        ws["A3"] = "Patient"
        ws["B3"] = "Date"
        ws["C3"] = "HbA1c"

        ws["A4"] = "ID*"
        ws["B4"] = "(dd-mmm-yyyy)"
        ws["C4"] = "%"

        ws["A5"] = "P001"  # Data starts here

        header_1, header_2 = read_header_rows(ws, data_start_row=5)

        assert header_1 == ["ID*", "(dd-mmm-yyyy)", "%"]
        assert header_2 == ["Patient", "Date", "HbA1c"]

        wb.close()

    def test_trims_to_last_non_none_column(self):
        """Test that headers are trimmed to last non-None column."""
        wb = Workbook()
        ws = wb.active

        # Data starts at row 10
        ws["A8"] = "Patient"
        ws["B8"] = "Name"
        ws["C8"] = "Age"
        # D8-Z8 remain None

        ws["A9"] = "ID*"
        ws["B9"] = None
        ws["C9"] = None

        ws["A10"] = "P001"

        header_1, header_2 = read_header_rows(ws, data_start_row=10)

        # Should trim to column C (last non-None)
        assert len(header_1) == 3
        assert len(header_2) == 3
        assert header_1 == ["ID*", None, None]
        assert header_2 == ["Patient", "Name", "Age"]

        wb.close()

    def test_realistic_tracker_width(self):
        """Test with realistic tracker dimensions (31 columns)."""
        wb = Workbook()
        ws = wb.active

        data_start_row = 77

        # Create 31 columns of headers
        for col_idx in range(1, 32):  # 1 to 31 inclusive
            ws.cell(row=75, column=col_idx, value=f"H2_Col{col_idx}")
            ws.cell(row=76, column=col_idx, value=f"H1_Col{col_idx}")

        # Put data at row 77
        ws.cell(row=77, column=1, value="P001")

        header_1, header_2 = read_header_rows(ws, data_start_row=data_start_row)

        assert len(header_1) == 31
        assert len(header_2) == 31
        assert header_1[0] == "H1_Col1"
        assert header_1[30] == "H1_Col31"
        assert header_2[0] == "H2_Col1"
        assert header_2[30] == "H2_Col31"

        wb.close()

    def test_mixed_none_values_in_headers(self):
        """Test headers with mixed None and non-None values."""
        wb = Workbook()
        ws = wb.active

        # Header row 2 (further from data)
        ws["A3"] = "Patient"
        ws["B3"] = None
        ws["C3"] = "Updated HbA1c"
        ws["D3"] = None  # Horizontally merged
        ws["E3"] = None

        # Header row 1 (closer to data)
        ws["A4"] = "ID*"
        ws["B4"] = "Name"
        ws["C4"] = "%"
        ws["D4"] = "(dd-mmm-yyyy)"
        ws["E4"] = None

        ws["A5"] = "P001"  # Data

        header_1, header_2 = read_header_rows(ws, data_start_row=5)

        # Should trim to column D (last non-None in header_1)
        assert len(header_1) == 4
        assert len(header_2) == 4
        assert header_1 == ["ID*", "Name", "%", "(dd-mmm-yyyy)"]
        assert header_2 == ["Patient", None, "Updated HbA1c", None]

        wb.close()

    def test_randomized_header_position(self):
        """Test with randomized data start position."""
        wb = Workbook()
        ws = wb.active

        # Random data start between rows 20 and 100
        random_data_start = random.randint(20, 100)
        header_row_1 = random_data_start - 1
        header_row_2 = random_data_start - 2

        # Set headers
        ws.cell(row=header_row_2, column=1, value="Header2")
        ws.cell(row=header_row_1, column=1, value="Header1")
        ws.cell(row=random_data_start, column=1, value="Data")

        header_1, header_2 = read_header_rows(ws, data_start_row=random_data_start)

        assert header_1 == ["Header1"]
        assert header_2 == ["Header2"]

        wb.close()

    def test_respects_max_cols_parameter(self):
        """Test that max_cols parameter limits the read width."""
        wb = Workbook()
        ws = wb.active

        # Create 200 columns of data
        for col_idx in range(1, 201):
            ws.cell(row=3, column=col_idx, value=f"H2_{col_idx}")
            ws.cell(row=4, column=col_idx, value=f"H1_{col_idx}")

        ws["A5"] = "Data"

        # Read with max_cols=50
        header_1, header_2 = read_header_rows(ws, data_start_row=5, max_cols=50)

        # Should only read up to column 50
        assert len(header_1) == 50
        assert len(header_2) == 50
        assert header_1[49] == "H1_50"

        wb.close()

    def test_all_none_headers(self):
        """Test when both header rows are completely None.

        Note: When no non-None values are found, the function returns
        max_cols None values (default behavior). In practice, this edge
        case doesn't occur as real trackers always have headers.
        """
        wb = Workbook()
        ws = wb.active

        # Headers are all None
        # (openpyxl cells are None by default)

        ws["A5"] = "Data"

        header_1, header_2 = read_header_rows(ws, data_start_row=5, max_cols=10)

        # Returns max_cols None values when nothing is found
        assert len(header_1) == 10
        assert len(header_2) == 10
        assert all(h is None for h in header_1)
        assert all(h is None for h in header_2)

        wb.close()


class TestMergeHeaders:
    """Tests for merge_headers() function."""

    def test_both_headers_present(self):
        """Test merging when both header rows have values."""
        h1 = ["%", "mmol/L", "kg"]
        h2 = ["HbA1c", "FBG", "Weight"]
        result = merge_headers(h1, h2)
        assert result == ["HbA1c %", "FBG mmol/L", "Weight kg"]

    def test_only_h2_present(self):
        """Test when only header row 2 has values."""
        h1 = [None, None, None]
        h2 = ["Patient ID", "Name", "Age"]
        result = merge_headers(h1, h2)
        assert result == ["Patient ID", "Name", "Age"]

    def test_only_h1_present(self):
        """Test when only header row 1 has values (single-line headers)."""
        h1 = ["Patient ID", "Name", "Age"]
        h2 = [None, None, None]
        result = merge_headers(h1, h2)
        assert result == ["Patient ID", "Name", "Age"]

    def test_horizontal_merge_forward_fill(self):
        """Test forward-fill with synonym validation.

        Forward-fill happens when mapper validates the combined header.
        """
        h1 = ["%", "(dd-mmm-yyyy)", "mmol/L", "(dd-mmm-yyyy)"]
        h2 = ["Updated HbA1c", None, "Updated FBG", None]
        # Mock mapper that knows these forward-filled patterns
        mapper = create_mock_mapper(
            {
                "Updated HbA1c %",
                "Updated HbA1c (dd-mmm-yyyy)",
                "Updated FBG mmol/L",
                "Updated FBG (dd-mmm-yyyy)",
            }
        )
        result = merge_headers(h1, h2, mapper)
        assert result == [
            "Updated HbA1c %",
            "Updated HbA1c (dd-mmm-yyyy)",
            "Updated FBG mmol/L",
            "Updated FBG (dd-mmm-yyyy)",
        ]

    def test_mixed_headers(self):
        """Test realistic mix of header patterns.

        Forward-fill happens when mapper validates the combined header.
        """
        h1 = ["ID*", "Name", "%", "(date)", None, "kg"]
        h2 = ["Patient", None, "HbA1c", None, "Notes", "Weight"]
        # Mock mapper that validates these forward-fills
        mapper = create_mock_mapper(
            {
                "Patient ID*",
                "Patient Name",
                "HbA1c %",
                "HbA1c (date)",
            }
        )
        result = merge_headers(h1, h2, mapper)
        assert result == [
            "Patient ID*",
            "Patient Name",  # Forward-filled and validated
            "HbA1c %",
            "HbA1c (date)",  # Forward-filled and validated
            "Notes",
            "Weight kg",
        ]

    def test_none_values_reset_forward_fill(self):
        """Test that None in both headers results in None.

        Forward-fill only happens when h1 exists and mapper validates.
        """
        h1 = ["%", "(date)", None, "kg"]
        h2 = ["HbA1c", None, None, "Weight"]
        # Mock mapper that validates HbA1c forward-fills
        mapper = create_mock_mapper(
            {
                "HbA1c %",
                "HbA1c (date)",
            }
        )
        result = merge_headers(h1, h2, mapper)
        assert result == [
            "HbA1c %",
            "HbA1c (date)",
            None,
            "Weight kg",
        ]

    def test_whitespace_normalization(self):
        """Test that extra whitespace and newlines are normalized."""
        h1 = ["ID\n(format)", "  Name  "]
        h2 = ["Patient\nID", "Full  Name"]
        result = merge_headers(h1, h2)
        assert result == [
            "Patient ID ID (format)",
            "Full Name Name",
        ]

    def test_updated_year_marker_inherits_previous_subject(self):
        """The 2022 template labels an update-date column "Updated 2022".

        The subject it belongs to is only in the column to its left, so
        "Updated <year>" alone must not become the column name. Without this,
        blood_pressure_updated and edu_occ_updated go unmapped on every 2022
        tracker -- 7,165 values.
        """
        h1 = ["mm HG", "Date"]
        h2 = ["Blood Pressure ", "Updated\n2022"]
        result = merge_headers(h1, h2)
        assert result == ["Blood Pressure mm HG", "Blood Pressure Date"]

    def test_updated_year_marker_on_patient_list(self):
        """Same marker, different subject: the Patient List's education column."""
        h1 = [None, "Date"]
        h2 = ["Level of Education\nOr Occupation", " Updated \n2022"]
        result = merge_headers(h1, h2)
        assert result == [
            "Level of Education Or Occupation",
            "Level of Education Or Occupation Date",
        ]

    def test_updated_year_marker_without_previous_subject(self):
        """With nothing to inherit, the marker is left alone rather than dropped."""
        h1 = ["Date"]
        h2 = ["Updated 2022"]
        result = merge_headers(h1, h2)
        assert result == ["Updated 2022 Date"]

    def test_updated_without_year_is_not_a_marker(self):
        """ "Updated HbA1c" is a real subject, not a continuation marker."""
        h1 = ["%"]
        h2 = ["Updated HbA1c"]
        result = merge_headers(h1, h2)
        assert result == ["Updated HbA1c %"]

    def test_empty_headers(self):
        """Test with empty header lists."""
        result = merge_headers([], [])
        assert result == []

    def test_single_column(self):
        """Test with single column."""
        h1 = ["ID"]
        h2 = ["Patient"]
        result = merge_headers(h1, h2)
        assert result == ["Patient ID"]


class TestMergeHeadersWithMergedSpans:
    """A merged upper header names every column its span covers.

    The 2021 Putrajaya layout: "Complication Screening (Current Month Testing)"
    is merged across five columns, so the sub-headers "Results" and
    "Date (mmm-yy)" further along the block have no upper header of their own.
    Forward-fill alone cannot reach them -- it resets at the blank columns
    between -- and the bare names map nowhere.
    """

    def test_merged_span_names_a_sub_header_beyond_a_blank_gap(self):
        h1 = ["Select for Drop Down", None, None, "Results", "Date (mmm-yy)"]
        h2 = ["Complication Screening", None, None, None, None]
        mapper = create_mock_mapper(
            {"Complication Screening Results", "Complication Screening Date (mmm-yy)"}
        )

        result = merge_headers(h1, h2, mapper=mapper, merged_spans=[(1, 5)])

        assert result[3] == "Complication Screening Results"
        assert result[4] == "Complication Screening Date (mmm-yy)"

    def test_a_column_with_no_sub_header_is_not_named_by_the_title(self):
        """A second column under one merged title is not the same field.

        The 2022 template merges "Insulin Regimen" across two columns whose
        second holds a near-duplicate of the first ("Basal-bolus MDI (AN/HI)"
        against "Basal-bolus (AN/HI)"). Naming both would comma-join them into
        one value, which is worse data than the first column alone.
        """
        h1 = [None, None]
        h2 = ["Insulin Regimen", None]

        result = merge_headers(h1, h2, merged_spans=[(1, 2)])

        assert result == ["Insulin Regimen", None]

    def test_a_multi_select_block_names_every_column_it_spans(self):
        """The 2022 complication-screening block records one screening per column.

        A patient screened for kidney, eye and foot in one month has three ticks
        in three columns under one merged title, and only the leftmost carries a
        header of its own. Giving the others the same header hands them to
        `merge_duplicate_columns_data`, which comma-joins them into one cell --
        measured across the corpus as 227 ticks in 133 patient-months that were
        otherwise dropped, none of them a repeat of the tick that was kept.
        """
        h1 = ["Drop Down", None, None]
        h2 = ["Current Month Complication Screening", None, None]
        mapper = create_mock_mapper(
            {"Current Month Complication Screening Drop Down"},
            {"Current Month Complication Screening Drop Down": "complication_screening"},
        )

        result = merge_headers(h1, h2, mapper=mapper, merged_spans=[(1, 3)])

        assert result == ["Current Month Complication Screening Drop Down"] * 3

    def test_a_multi_select_block_leaves_named_columns_alone(self):
        """Only the columns with no header of their own join the block."""
        h1 = ["Drop Down", None, "Results"]
        h2 = ["Current Month Complication Screening", None, None]
        mapper = create_mock_mapper(
            {
                "Current Month Complication Screening Drop Down",
                "Current Month Complication Screening Results",
            },
            {"Current Month Complication Screening Drop Down": "complication_screening"},
        )

        result = merge_headers(h1, h2, mapper=mapper, merged_spans=[(1, 3)])

        assert result == [
            "Current Month Complication Screening Drop Down",
            "Current Month Complication Screening Drop Down",
            "Current Month Complication Screening Results",
        ]

    def test_a_single_select_block_does_not_take_its_spanned_columns(self):
        """Insulin regimen is one value per patient-month, so the shadow stays out.

        2,910 of the 3,659 values under the 2022 insulin merge are byte-identical
        to the column the merge anchors; joining them would produce a worse value.
        """
        h1 = ["Drop Down", None]
        h2 = ["Insulin Regimen", None]
        mapper = create_mock_mapper(
            {"Insulin Regimen Drop Down"},
            {"Insulin Regimen Drop Down": "insulin_regimen"},
        )

        result = merge_headers(h1, h2, mapper=mapper, merged_spans=[(1, 2)])

        assert result == ["Insulin Regimen Drop Down", None]

    def test_two_sub_headers_qualifying_alike_do_not_merge(self):
        h1 = ["Date", None, "Date"]
        h2 = ["Screening", None, None]
        mapper = create_mock_mapper({"Screening Date"})

        result = merge_headers(h1, h2, mapper=mapper, merged_spans=[(1, 3)])

        assert result == ["Screening Date", None, "Date"]

    def test_without_spans_behaviour_is_unchanged(self):
        h1 = ["Select for Drop Down", None, "Results"]
        h2 = ["Complication Screening", None, None]

        assert merge_headers(h1, h2) == [
            "Complication Screening Select for Drop Down",
            None,
            "Results",
        ]


class TestRecoverBlankHeadersInsideMergedSpan:
    """A merged header outranks a sibling sheet's guess.

    Putrajaya's month sheets do not share one layout: Jul21 has "Patient
    Observations" at the position Dec21 uses for a complication-screening
    selection. Recovering by position files a screening result under
    observations. The merged span is the workbook's own statement about what
    the column belongs to, so it wins.
    """

    def test_abstains_for_a_position_inside_a_merged_span(self):
        headers = ["ID", "Complication Screening", None]
        data = [("1", "Kidneys", "Foot Examination (Nerves)")]
        siblings = [["ID", "Complication Screening", "Patient Observations"]]

        result = recover_blank_headers(headers, data, siblings, merged_spans=[(2, 3)])

        assert result == ["ID", "Complication Screening", None]

    def test_recovery_outside_any_span_still_works(self):
        """2021 Kantha Bopha's insulin-regimen recovery must survive."""
        headers = ["ID", None]
        data = [("1", "Self-mixed BD")]
        siblings = [["ID", "Insulin Regime"]]

        result = recover_blank_headers(headers, data, siblings, merged_spans=[(5, 7)])

        assert result == ["ID", "Insulin Regime"]


class TestFilterValidColumns:
    """Tests for filter_valid_columns() function."""

    def test_all_valid_headers(self):
        """Test when all headers are valid (no None)."""
        headers = ["ID", "Name", "Age"]
        data = [("1", "Alice", "30"), ("2", "Bob", "25")]
        valid_headers, filtered_data = filter_valid_columns(headers, data)

        assert valid_headers == ["ID", "Name", "Age"]
        assert filtered_data == [["1", "Alice", "30"], ["2", "Bob", "25"]]

    def test_some_none_headers(self):
        """Test filtering out None headers."""
        headers = ["ID", None, "Name", None, "Age"]
        data = [("1", "x", "Alice", "y", "30"), ("2", "x", "Bob", "y", "25")]
        valid_headers, filtered_data = filter_valid_columns(headers, data)

        assert valid_headers == ["ID", "Name", "Age"]
        assert filtered_data == [["1", "Alice", "30"], ["2", "Bob", "25"]]

    def test_all_none_headers(self):
        """Test when all headers are None."""
        headers = [None, None, None]
        data = [("1", "2", "3"), ("4", "5", "6")]
        valid_headers, filtered_data = filter_valid_columns(headers, data)

        assert valid_headers == []
        assert filtered_data == []

    def test_empty_data(self):
        """Test with empty data."""
        headers = ["ID", "Name"]
        data = []
        valid_headers, filtered_data = filter_valid_columns(headers, data)

        assert valid_headers == ["ID", "Name"]
        assert filtered_data == []

    def test_single_valid_column(self):
        """Test with single valid column."""
        headers = [None, "ID", None]
        data = [("x", "1", "y"), ("x", "2", "y")]
        valid_headers, filtered_data = filter_valid_columns(headers, data)

        assert valid_headers == ["ID"]
        assert filtered_data == [["1"], ["2"]]

    def test_preserves_order(self):
        """Test that column order is preserved."""
        headers = ["A", None, "B", None, "C", "D", None]
        data = [(1, 2, 3, 4, 5, 6, 7)]
        valid_headers, filtered_data = filter_valid_columns(headers, data)

        assert valid_headers == ["A", "B", "C", "D"]
        assert filtered_data == [[1, 3, 5, 6]]


class TestFindDroppedDataColumns:
    """Tests for find_dropped_data_columns() function."""

    def test_reports_headerless_column_holding_text(self):
        headers = ["ID", None, "Age"]
        data = [("1", "Self-mixed BD", "30"), ("2", "Basal-Bolus", "25")]

        assert find_dropped_data_columns(headers, data) == [(1, 2)]

    def test_ignores_headerless_column_that_is_empty(self):
        headers = ["ID", None, "Age"]
        data = [("1", None, "30"), ("2", "", "25")]

        assert find_dropped_data_columns(headers, data) == []

    def test_ignores_row_counter(self):
        """A headerless all-numeric column is the tracker's own row counter."""
        headers = [None, "ID"]
        data = [("1", "KH_PK001"), ("2", "KH_PK002"), ("3.0", "KH_PK003")]

        assert find_dropped_data_columns(headers, data) == []

    def test_reports_mixed_numeric_and_text_column(self):
        """A stray number among real text does not make the column a counter."""
        headers = [None, "ID"]
        data = [("7", "KH_JV001"), ("*Needs Meter", "KH_JV002")]

        assert find_dropped_data_columns(headers, data) == [(0, 2)]

    def test_ignores_columns_that_have_headers(self):
        headers = ["ID", "Notes"]
        data = [("1", "text")]

        assert find_dropped_data_columns(headers, data) == []

    def test_ignores_the_second_half_of_a_merged_heading(self):
        """A heading stretched over two columns names both of them.

        The 2022 template merges "Insulin Regimen" across a pair of columns, so
        the right-hand one reads as headerless while the value is read from the
        left-hand one. Reporting it tells the clinic to fix a heading that is
        not broken.
        """
        headers = ["ID", "Insulin Regimen", None]
        data = [("1", "Self-mixed BD", "Self-mixed BD")]

        assert find_dropped_data_columns(headers, data, merged_spans=[(2, 3)]) == []

    def test_reports_a_headerless_column_outside_any_merge(self):
        headers = ["ID", "Insulin Regimen", None]
        data = [("1", "Self-mixed BD", "Mixtard30 Penfill")]

        assert find_dropped_data_columns(headers, data, merged_spans=[(2, 2)]) == [(2, 1)]

    def test_ignores_a_row_counter_carrying_one_stray_keystroke(self):
        """2023 Mahosot has 83 row numbers and one cell reading "m"."""
        headers = [None, "ID"]
        data = [("1", "a"), ("2", "b"), ("m", "c"), ("4", "d")]

        assert find_dropped_data_columns(headers, data) == []

    def test_still_ignores_a_counter_that_restarts(self):
        """The ascending test is only asked of a column carrying a stray."""
        headers = [None, "ID"]
        data = [("1", "a"), ("2", "b"), ("1", "c")]

        assert find_dropped_data_columns(headers, data) == []

    def test_reports_a_column_of_numbers_with_a_stray_that_do_not_count(self):
        headers = [None, "ID"]
        data = [("70", "a"), ("65", "b"), ("high", "c")]

        assert find_dropped_data_columns(headers, data) == [(0, 3)]

    def test_treats_a_whitespace_only_cell_as_empty(self):
        headers = ["ID", None]
        data = [("1", "   "), ("2", None)]

        assert find_dropped_data_columns(headers, data) == []


class TestRecoverBlankHeaders:
    """Tests for recover_blank_headers() function."""

    def test_recovers_from_the_one_sibling_that_names_the_column(self):
        headers = ["ID", None]
        data = [("1", "Self-mixed BD")]
        siblings = [["ID", "Insulin Regime"]]

        assert recover_blank_headers(headers, data, siblings) == ["ID", "Insulin Regime"]

    def test_abstains_when_siblings_disagree(self):
        """Two candidate names is a guess, not a recovery."""
        headers = ["ID", None]
        data = [("1", "x")]
        siblings = [["ID", "Complication Screening"], ["ID", "Hospitalisation Date"]]

        assert recover_blank_headers(headers, data, siblings) == ["ID", None]

    def test_agreeing_siblings_are_not_a_disagreement(self):
        headers = ["ID", None]
        data = [("1", "x")]
        siblings = [["ID", "Clinic Visit"], ["ID", "Clinic Visit"], ["ID", None]]

        assert recover_blank_headers(headers, data, siblings) == ["ID", "Clinic Visit"]

    def test_abstains_when_no_sibling_names_the_column(self):
        """The 2022 template's hidden merged-cell column is blank in every sheet."""
        headers = ["ID", None]
        data = [("1", "Self-mixed BD")]
        siblings = [["ID", None], ["ID", None]]

        assert recover_blank_headers(headers, data, siblings) == ["ID", None]

    def test_leaves_an_empty_column_alone(self):
        """A blank header over no data is a spacer, not a lost column."""
        headers = ["ID", None]
        data = [("1", None)]
        siblings = [["ID", "Insulin Regime"]]

        assert recover_blank_headers(headers, data, siblings) == ["ID", None]

    def test_ignores_the_row_counter(self):
        headers = [None, "ID"]
        data = [("1", "KH_PK001"), ("2", "KH_PK002")]
        siblings = [["Nr", "ID"]]

        assert recover_blank_headers(headers, data, siblings) == [None, "ID"]

    def test_never_creates_a_duplicate_column_name(self):
        """A donor name this sheet already uses would collide on rename."""
        headers = ["ID", "Insulin Regime", None]
        data = [("1", "a", "b")]
        siblings = [["ID", "Insulin Regime", "Insulin Regime"]]

        assert recover_blank_headers(headers, data, siblings) == ["ID", "Insulin Regime", None]

    def test_tolerates_siblings_of_different_width(self):
        headers = ["ID", None]
        data = [("1", "x")]
        siblings = [["ID"], ["ID", "Clinic Visit"]]

        assert recover_blank_headers(headers, data, siblings) == ["ID", "Clinic Visit"]

    def test_no_siblings_is_a_no_op(self):
        headers = ["ID", None]
        data = [("1", "x")]

        assert recover_blank_headers(headers, data, []) == ["ID", None]


class TestFindLayoutChanges:
    """Tests for find_layout_changes() function."""

    def test_flags_a_position_whose_canonical_column_changes(self):
        layouts = {
            "Jan20": ["Patient ID", "Baseline FBG (mmol/dL)"],
            "Jun20": ["Patient ID", "Baseline FBG (mg/dL)"],
        }
        mapper = create_mock_mapper(set())
        mapper.get_standard_name = lambda h: {
            "Baseline FBG (mmol/dL)": "fbg_baseline_mmol",
            "Baseline FBG (mg/dL)": "fbg_baseline_mg",
        }.get(h, h)

        changes = find_layout_changes(layouts, mapper)

        assert len(changes) == 1
        assert changes[0].index == 1
        assert changes[0].renames_only is False
        assert changes[0].headers == ["Baseline FBG (mg/dL)", "Baseline FBG (mmol/dL)"]

    def test_a_pure_rename_is_flagged_but_marked_harmless(self):
        layouts = {
            "Jan20": ["Patient ID", "Insulin regime"],
            "Mar20": ["Patient ID", "Insulin regimen"],
        }
        mapper = create_mock_mapper(set())
        mapper.get_standard_name = lambda h: "insulin_regimen" if "nsulin" in h else h

        changes = find_layout_changes(layouts, mapper)

        assert len(changes) == 1
        assert changes[0].renames_only is True

    def test_sheets_that_agree_produce_nothing(self):
        layouts = {"Jan20": ["Patient ID", "Age"], "Feb20": ["Patient ID", "Age"]}
        mapper = create_mock_mapper(set())
        mapper.get_standard_name = lambda h: h

        assert find_layout_changes(layouts, mapper) == []

    def test_a_blank_header_is_not_a_change(self):
        """One sheet leaving a header empty is recover_blank_headers' business."""
        layouts = {"Jan20": ["Patient ID", "Age"], "Feb20": ["Patient ID", None]}
        mapper = create_mock_mapper(set())
        mapper.get_standard_name = lambda h: h

        assert find_layout_changes(layouts, mapper) == []

    def test_a_single_sheet_cannot_disagree_with_itself(self):
        mapper = create_mock_mapper(set())
        mapper.get_standard_name = lambda h: h

        assert find_layout_changes({"Jan20": ["Patient ID"]}, mapper) == []

    def test_sheets_of_different_width_are_compared_where_they_overlap(self):
        layouts = {"Jan20": ["ID", "Insulin Regimen", "BASAL"], "Apr20": ["ID", "BASAL"]}
        mapper = create_mock_mapper(set())
        mapper.get_standard_name = lambda h: h

        changes = find_layout_changes(layouts, mapper)

        assert [c.index for c in changes] == [1]


class TestJoinStaticSheet:
    """The whole-tracker joins key on the normalized ID, not the raw one (ticket 58)."""

    @staticmethod
    def _monthly(ids: list[str]) -> pl.DataFrame:
        return pl.DataFrame({"patient_id": ids, "fbg_updated_mg": [1.0] * len(ids)})

    @staticmethod
    def _static(ids: list[str]) -> pl.DataFrame:
        return pl.DataFrame(
            {"patient_id": ids, "dob": [f"2010-01-0{i + 1}" for i in range(len(ids))]}
        )

    @classmethod
    def _join(cls, monthly_ids: list[str], static_ids: list[str]) -> pl.DataFrame:
        return join_static_sheet(
            cls._monthly(monthly_ids), cls._static(static_ids), ".static", "Patient List"
        )

    def test_hyphen_spelled_monthly_id_finds_its_underscore_patient_list_entry(self):
        """2023/2024 Mahosot: 672 rows lost every demographic to this mismatch."""
        assert self._join(["LA-MH056"], ["LA_MH056"])["dob"].to_list() == ["2010-01-01"]

    def test_transfer_clinic_suffix_does_not_block_the_match(self):
        assert self._join(["MY_SM003_SB"], ["MY_SM003"])["dob"].to_list() == ["2010-01-01"]

    def test_raw_patient_id_column_is_not_rewritten(self):
        """The comparison's row-alignment key depends on the raw spelling surviving."""
        assert self._join(["LA-MH056"], ["LA_MH056"])["patient_id"].to_list() == ["LA-MH056"]

    def test_no_join_key_column_leaks_into_the_output(self):
        result = self._join(["LA-MH056"], ["LA_MH056"])
        assert result.columns == ["patient_id", "fbg_updated_mg", "dob"]

    def test_unmatched_monthly_id_keeps_its_row_with_null_demographics(self):
        """MM_YG013_MG has no Patient List entry at all -- a source defect, not a key bug."""
        result = self._join(["MM_YG013_MG"], ["MM_MG013"])
        assert result.height == 1
        assert result["dob"].to_list() == [None]

    def test_a_row_never_gains_a_second_partner(self):
        """Two spellings folding to one key must not duplicate the month row."""
        assert self._join(["TH_ST029"], ["TH-ST029", "TH_ST029"]).height == 1

    def test_colliding_columns_take_the_suffix_the_caller_asked_for(self):
        monthly = pl.DataFrame({"patient_id": ["LA-MH056"], "province": ["monthly"]})
        static = pl.DataFrame({"patient_id": ["LA_MH056"], "province": ["static"]})
        result = join_static_sheet(monthly, static, ".static", "Patient List")
        assert result["province"].to_list() == ["monthly"]
        assert result["province.static"].to_list() == ["static"]

    def test_the_annual_sheet_join_gets_the_same_key_and_its_own_suffix(self):
        """2024 CDA and 2026 Surat Thani lose Annual columns to the same mismatch."""
        annual = pl.DataFrame({"patient_id": ["KH_CD016"], "edu_occ": ["Student"]})
        result = join_static_sheet(self._monthly(["KH-CD016"]), annual, ".annual", "Annual")
        assert result["edu_occ"].to_list() == ["Student"]


class TestReportDuplicatePatientsInSheet:
    """One patient written down twice on one month sheet.

    Vietnam National Children's ``Jul24`` splices two lists into one sheet: the
    row-number column runs 1..78 while 21 patient IDs repeat, each copy holding
    different readings. Both copies reach the monthly table, so that clinic's
    July is counted twice for those patients.
    """

    def test_reports_a_patient_listed_twice(self, collector):
        df = pl.DataFrame({"patient_id": ["VN_VC002", "VN_VC003", "VN_VC002"]})

        report_duplicate_patients_in_sheet(df, "Jul24")

        reported = collector.to_dataframe()
        assert reported.height == 1
        row = reported.row(0, named=True)
        assert row["error_code"] == "duplicate_patient_row_in_sheet"
        assert row["patient_id"] == "VN_VC002"
        assert row["sheet_name"] == "Jul24"
        assert "2 rows" in row["message"]

    def test_reports_each_duplicated_patient_once(self, collector):
        df = pl.DataFrame({"patient_id": ["A_B001", "A_B002", "A_B001", "A_B002", "A_B003"]})

        report_duplicate_patients_in_sheet(df, "Jul24")

        assert sorted(collector.to_dataframe()["patient_id"].to_list()) == ["A_B001", "A_B002"]

    def test_counts_every_copy(self, collector):
        df = pl.DataFrame({"patient_id": ["A_B001"] * 3})

        report_duplicate_patients_in_sheet(df, "Jul24")

        assert "3 rows" in collector.to_dataframe().row(0, named=True)["message"]

    def test_two_spellings_of_one_id_are_one_patient(self, collector):
        """Surat Thani writes ``TH-ST029`` and ``TH_ST029`` on the same sheet.

        The two fold to one identity downstream, so the patient ends up with
        two rows for one month; the message names both spellings, since the
        workbook fix is to make them agree.
        """
        df = pl.DataFrame({"patient_id": ["TH-ST029", "TH_ST029"]})

        report_duplicate_patients_in_sheet(df, "May26")

        row = collector.to_dataframe().row(0, named=True)
        assert row["patient_id"] == "TH_ST029"
        assert "TH-ST029" in row["original_value"]
        assert "TH_ST029" in row["original_value"]

    def test_silent_when_every_patient_appears_once(self, collector):
        df = pl.DataFrame({"patient_id": ["A_B001", "A_B002", "A_B003"]})

        report_duplicate_patients_in_sheet(df, "Jul24")

        assert len(collector) == 0

    def test_rows_with_no_id_are_not_duplicates(self, collector):
        """An unidentifiable row is a different defect, already reported."""
        df = pl.DataFrame({"patient_id": [None, None, "A_B001"]})

        report_duplicate_patients_in_sheet(df, "Jul24")

        assert len(collector) == 0

    def test_broken_formulas_are_not_one_duplicated_patient(self, collector):
        """``#REF!`` on nine rows is nine unidentifiable patients, not one.

        Those rows are dropped and reported as ``excel_error_patient_id``
        further down; calling them a duplicate would say the opposite of what
        is wrong with them.
        """
        df = pl.DataFrame({"patient_id": ["#REF!", "#REF!", "#REF!"]})

        report_duplicate_patients_in_sheet(df, "May26")

        assert len(collector) == 0

    def test_numeric_filler_ids_are_not_a_patient(self, collector):
        """Template rows left with ``0`` in the ID cell name no patient."""
        df = pl.DataFrame({"patient_id": ["0", "0", "0.0", "0.0"]})

        report_duplicate_patients_in_sheet(df, "Jun26")

        assert len(collector) == 0
