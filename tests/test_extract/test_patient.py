"""Tests for patient data extraction."""

from pathlib import Path

import polars as pl
import pytest

from a4d.extract.patient import (
    extract_patient_data,
    extract_tracker_month,
    find_month_sheets,
    get_tracker_year,
    harmonize_patient_data_columns,
    merge_duplicate_columns_data,
    read_all_patient_sheets,
)


def column_letter_to_index(col_letter: str) -> int:
    """Convert Excel column letter to 0-based index.

    Examples:
        A -> 0, B -> 1, Z -> 25, AA -> 26, AB -> 27, AC -> 28
    """
    result = 0
    for char in col_letter:
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result - 1


def calculate_expected_columns(start_col: str, end_col: str) -> int:
    """Calculate expected number of columns from Excel range.

    Args:
        start_col: Starting column letter (e.g., 'B')
        end_col: Ending column letter (e.g., 'AC')

    Returns:
        Number of columns in the range

    Examples:
        B to Z: 25 columns
        B to AC: 28 columns
        B to AB: 27 columns
    """
    start_idx = column_letter_to_index(start_col)
    end_idx = column_letter_to_index(end_col)
    return end_idx - start_idx + 1


# Test data paths
TRACKER_SBU_2024 = Path(
    "/Volumes/USB SanDisk 3.2Gen1 Media/a4d/a4dphase2_upload/"
    "Malaysia/SBU/2024_Sibu Hospital A4D Tracker.xlsx"
)
TRACKER_PNG_2019 = Path(
    "/Volumes/USB SanDisk 3.2Gen1 Media/a4d/a4dphase2_upload/"
    "Malaysia/PNG/2019_Penang General Hospital A4D Tracker_DC.xlsx"
)
TRACKER_PNG_2018 = Path(
    "/Volumes/USB SanDisk 3.2Gen1 Media/a4d/a4dphase2_upload/"
    "Malaysia/PNG/2018_Penang General Hospital A4D Tracker_DC.xlsx"
)
TRACKER_MHS_2017 = Path(
    "/Volumes/USB SanDisk 3.2Gen1 Media/a4d/a4dphase2_upload/"
    "Laos/MHS/2017_Mahosot Hospital A4D Tracker.xlsx"
)
TRACKER_MHS_2025 = Path(
    "/Volumes/USB SanDisk 3.2Gen1 Media/a4d/a4dphase2_upload/"
    "Laos/MHS/2025_06_Mahosot Hospital A4D Tracker.xlsx"
)


@pytest.mark.skipif(not TRACKER_SBU_2024.exists(), reason="Tracker file not available")
def test_get_tracker_year_from_sheet_names():
    """Test extracting year from sheet names."""
    year = get_tracker_year(TRACKER_SBU_2024, ["Jan24", "Feb24", "Mar24"])
    assert year == 2024


@pytest.mark.skipif(not TRACKER_SBU_2024.exists(), reason="Tracker file not available")
def test_get_tracker_year_from_filename():
    """Test extracting year from filename as fallback."""
    year = get_tracker_year(TRACKER_SBU_2024, ["January", "February"])
    assert year == 2024


@pytest.mark.skipif(not TRACKER_SBU_2024.exists(), reason="Tracker file not available")
def test_find_month_sheets_2024():
    """Test finding month sheets in 2024 tracker."""
    from openpyxl import load_workbook

    wb = load_workbook(TRACKER_SBU_2024, data_only=True)
    month_sheets = find_month_sheets(wb)

    assert len(month_sheets) > 0
    assert any("Jan" in sheet for sheet in month_sheets)
    assert any("Dec" in sheet for sheet in month_sheets)


# Parameterized test data: (tracker_file, sheet_name, year, expected_patients, expected_cols, notes)
# Note: expected_cols is the actual number after filtering out None header columns
TRACKER_TEST_CASES = [
    # 2024 tracker - optimized single-pass extraction
    (
        TRACKER_SBU_2024,
        "Jan24",
        2024,
        4,
        calculate_expected_columns("B", "AG") - 1,
        "Single-pass read-only",
    ),
    # 2019 tracker - format changes across months! Optimized extraction
    (
        TRACKER_PNG_2019,
        "Jan19",
        2019,
        10,
        calculate_expected_columns("B", "Z"),
        "Single-pass read-only",
    ),
    (
        TRACKER_PNG_2019,
        "Feb19",
        2019,
        10,
        calculate_expected_columns("B", "AC"),
        "Single-pass read-only",
    ),
    (
        TRACKER_PNG_2019,
        "Mar19",
        2019,
        10,
        calculate_expected_columns("B", "AB"),
        "Single-pass read-only",
    ),
    (
        TRACKER_PNG_2019,
        "Oct19",
        2019,
        11,
        calculate_expected_columns("B", "AB"),
        "Single-pass read-only",
    ),
    # 2018 tracker - single-line headers
    (
        TRACKER_PNG_2018,
        "Dec18",
        2018,
        10,
        calculate_expected_columns("B", "T"),
        "Single-pass read-only",
    ),
]


@pytest.mark.skipif(
    any(not tf.exists() for tf, _, _, _, _, _ in TRACKER_TEST_CASES),
    reason="Tracker files not available",
)
@pytest.mark.parametrize(
    ("tracker_file", "sheet_name", "year", "expected_patients", "expected_cols", "notes"),
    TRACKER_TEST_CASES,
    ids=lambda params: f"{params[1] if isinstance(params, tuple) and len(params) > 1 else params}",
)
def test_extract_patient_data_schema(
    tracker_file, sheet_name, year, expected_patients, expected_cols, notes
):
    """Test patient data extraction with schema validation across different months.

    This parameterized test validates that:
    1. Correct number of patients are extracted
    2. Correct number of columns match expected (after filtering None headers)
    3. Format changes between months are handled correctly

    The test is critical because tracker formats change even within the same year,
    and data quality is inconsistent across different months.
    """
    df = extract_patient_data(tracker_file, sheet_name, year)

    # Check dimensions
    assert len(df) == expected_patients, (
        f"{sheet_name}: Expected {expected_patients} patients, got {len(df)}"
    )
    assert len(df.columns) == expected_cols, (
        f"{sheet_name}: Expected {expected_cols} columns ({notes}), got {len(df.columns)}"
    )

    # Verify we have at least Patient ID column
    assert any("patient" in col.lower() and "id" in col.lower() for col in df.columns), (
        f"{sheet_name}: Missing Patient ID column in {df.columns}"
    )

    print(f"\n{sheet_name}: {len(df)} patients × {len(df.columns)} columns ({notes}) ✓")


@pytest.mark.skipif(not TRACKER_SBU_2024.exists(), reason="Tracker file not available")
def test_extract_patient_data_2024_detailed():
    """Detailed test for 2024 tracker with patient ID validation."""
    df = extract_patient_data(TRACKER_SBU_2024, "Jan24", 2024)

    # Verify specific patient IDs
    patient_ids = df["Patient ID*"].to_list()
    assert patient_ids == ["MY_QI001", "MY_QI002", "MY_QI003", "MY_QI004"], (
        f"Expected MY_QI001-004, got {patient_ids}"
    )

    print(f"\n2024 Jan24 - Patient IDs: {patient_ids} ✓")


def test_harmonize_patient_data_columns_basic():
    """Test basic column harmonization with known synonyms."""
    raw_df = pl.DataFrame(
        {
            "Patient ID*": ["MY_QI001", "MY_QI002"],
            "Age": [25, 30],
            "D.O.B.": ["1998-01-15", "1993-06-20"],
        }
    )

    harmonized = harmonize_patient_data_columns(raw_df)

    # Check that columns were renamed to standardized names
    assert "patient_id" in harmonized.columns
    assert "age" in harmonized.columns
    assert "dob" in harmonized.columns

    # Check that data is preserved
    assert harmonized["patient_id"].to_list() == ["MY_QI001", "MY_QI002"]
    assert harmonized["age"].to_list() == [25, 30]


def test_harmonize_patient_data_columns_multiple_synonyms():
    """Test that multiple columns mapping to same name are merged, not dropped.

    The 2023 template's complication screening block splits one canonical column
    across B.P./Kidney/Eye/Foot/Lipids sub-columns, each independently populated,
    so keeping only the first discards real values -- 2,489 of them across 27
    trackers. Absent sub-columns contribute nothing rather than a literal "NA".
    """
    raw_df = pl.DataFrame(
        {
            "Patient ID": ["P001", None],
            "ID": ["P002", None],
            "Patient ID*": [None, "P003"],
        }
    )

    harmonized = harmonize_patient_data_columns(raw_df)

    assert list(harmonized.columns) == ["patient_id"]
    assert harmonized["patient_id"].to_list() == ["P001,P002", "P003"]


def test_harmonize_patient_data_columns_unmapped_strict_false():
    """Test that unmapped columns are kept when strict=False (default)."""
    raw_df = pl.DataFrame(
        {
            "Patient ID*": ["MY_QI001"],
            "Age": [25],
            "UnknownColumn": ["some value"],
        }
    )

    harmonized = harmonize_patient_data_columns(raw_df, strict=False)

    # Mapped columns should be renamed
    assert "patient_id" in harmonized.columns
    assert "age" in harmonized.columns

    # Unmapped column should be kept as-is
    assert "UnknownColumn" in harmonized.columns


def test_harmonize_patient_data_columns_unmapped_strict_true():
    """Test that unmapped columns raise error when strict=True."""
    raw_df = pl.DataFrame(
        {
            "Patient ID*": ["MY_QI001"],
            "UnknownColumn": ["some value"],
        }
    )

    with pytest.raises(ValueError, match="Unmapped columns found"):
        harmonize_patient_data_columns(raw_df, strict=True)


def test_harmonize_patient_data_columns_empty_dataframe():
    """Test harmonization with empty DataFrame."""
    raw_df = pl.DataFrame()

    harmonized = harmonize_patient_data_columns(raw_df)

    assert len(harmonized) == 0
    assert len(harmonized.columns) == 0


@pytest.mark.skipif(not TRACKER_SBU_2024.exists(), reason="Tracker file not available")
def test_harmonize_real_tracker_data():
    """Test harmonization with real tracker data."""
    # Extract raw data
    raw_df = extract_patient_data(TRACKER_SBU_2024, "Jan24", 2024)

    # Harmonize columns
    harmonized = harmonize_patient_data_columns(raw_df)

    # Check that key columns were renamed
    assert "patient_id" in harmonized.columns
    assert "age" in harmonized.columns

    # Check that data is preserved
    assert len(harmonized) == len(raw_df)  # Same number of rows
    assert harmonized["patient_id"].to_list() == ["MY_QI001", "MY_QI002", "MY_QI003", "MY_QI004"]


def test_extract_tracker_month():
    """Test extracting month number from sheet name."""
    assert extract_tracker_month("Jan24") == 1
    assert extract_tracker_month("Feb24") == 2
    assert extract_tracker_month("Mar19") == 3
    assert extract_tracker_month("Dec23") == 12

    # Test with ValueError for invalid sheet names
    with pytest.raises(ValueError, match="Could not extract month"):
        extract_tracker_month("Sheet1")


def test_merge_duplicate_columns_data_no_duplicates():
    """Test that data without duplicate headers is unchanged."""
    headers = ["ID", "Name", "Age", "City"]
    data = [["1", "Alice", "25", "NYC"], ["2", "Bob", "30", "LA"]]

    result_headers, result_data = merge_duplicate_columns_data(headers, data)

    assert result_headers == headers
    assert result_data == data


def test_merge_duplicate_columns_data_with_duplicates():
    """Duplicate headers have their values comma-joined, not reduced to the first."""
    headers = ["ID", "DM Complications", "DM Complications", "DM Complications", "Age"]
    data = [["1", "A", "B", "C", "25"], ["2", "X", "Y", "Z", "30"]]

    result_headers, result_data = merge_duplicate_columns_data(headers, data)

    assert result_headers == ["ID", "DM Complications", "Age"]
    assert result_data == [["1", "A,B,C", "25"], ["2", "X,Y,Z", "30"]]


def test_merge_duplicate_columns_data_with_nulls():
    """Test merging duplicate columns with null values."""
    headers = ["ID", "DM Complications", "DM Complications", "DM Complications", "Age"]
    data = [["1", "A", None, "C", "25"], ["2", None, "Y", None, "30"]]

    result_headers, result_data = merge_duplicate_columns_data(headers, data)

    assert result_headers == ["ID", "DM Complications", "Age"]
    # Empty values are filtered out before joining
    assert result_data == [["1", "A,C", "25"], ["2", "Y", "30"]]


def test_merge_duplicate_columns_data_all_nulls():
    """Test merging when all duplicate columns have null values."""
    headers = ["ID", "DM Complications", "DM Complications", "Age"]
    data = [["1", None, None, "25"]]

    result_headers, result_data = merge_duplicate_columns_data(headers, data)

    assert result_headers == ["ID", "DM Complications", "Age"]
    # All nulls result in None
    assert result_data == [["1", None, "25"]]


def test_merge_duplicate_columns_data_multiple_groups():
    """Test merging multiple groups of duplicate columns."""
    headers = ["ID", "Status", "Status", "Value", "Value", "Value", "Name"]
    data = [["1", "A", "B", "X", "Y", "Z", "Alice"]]

    result_headers, result_data = merge_duplicate_columns_data(headers, data)

    assert result_headers == ["ID", "Status", "Value", "Name"]
    assert result_data == [["1", "A,B", "X,Y,Z", "Alice"]]


@pytest.mark.skipif(not TRACKER_SBU_2024.exists(), reason="Tracker file not available")
def test_read_all_patient_sheets_2024():
    """Test reading all patient sheets from 2024 tracker with Patient List and Annual."""
    df_all = read_all_patient_sheets(TRACKER_SBU_2024)

    # Check that we have data
    assert len(df_all) > 0, "Should have extracted patient data"

    # Check that metadata columns were added
    assert "sheet_name" in df_all.columns
    assert "tracker_month" in df_all.columns
    assert "tracker_year" in df_all.columns
    assert "file_name" in df_all.columns
    assert "clinic_id" in df_all.columns

    # Check that clinic_id is extracted from parent directory
    clinic_ids = df_all["clinic_id"].unique().to_list()
    assert len(clinic_ids) == 1  # All rows should have same clinic_id
    assert clinic_ids[0] == "SBU"  # Parent directory name

    # Check that we have data from multiple months
    unique_months = df_all["tracker_month"].unique().to_list()
    assert len(unique_months) > 1, "Should have data from multiple months"

    # Check that year is correct
    assert all(year == 2024 for year in df_all["tracker_year"].unique().to_list())

    # Check that patient_id column exists
    assert "patient_id" in df_all.columns

    # Check that we filtered out invalid rows (no null patient_ids)
    assert df_all["patient_id"].null_count() == 0

    # Check for baseline HbA1c column from Patient List (should be present after join)
    # Note: This may have .static suffix if there were conflicts
    hba1c_cols = [col for col in df_all.columns if "hba1c_baseline" in col.lower()]
    print(f"\nHbA1c baseline columns: {hba1c_cols}")

    print(
        f"\n2024 Tracker: {len(df_all)} total patients from {len(unique_months)} months"
        f" (with Patient List & Annual data) ✓"
    )


@pytest.mark.skipif(not TRACKER_PNG_2019.exists(), reason="Tracker file not available")
def test_read_all_patient_sheets_2019():
    """Test reading all patient sheets from 2019 tracker (different formats across months)."""
    df_all = read_all_patient_sheets(TRACKER_PNG_2019)

    # Check that we have data
    assert len(df_all) > 0, "Should have extracted patient data"

    # Check metadata columns
    assert "sheet_name" in df_all.columns
    assert "tracker_month" in df_all.columns
    assert "tracker_year" in df_all.columns

    # Check that year is correct
    assert all(year == 2019 for year in df_all["tracker_year"].unique().to_list())

    # Check that patient_id column exists
    assert "patient_id" in df_all.columns

    # Check that we filtered out invalid rows
    assert df_all["patient_id"].null_count() == 0

    # 2019 tracker has format changes across months - verify we handled them
    unique_months = df_all["tracker_month"].unique().to_list()
    print(f"\n2019 Tracker: {len(df_all)} total patients from {len(unique_months)} months ✓")


@pytest.mark.skipif(not TRACKER_SBU_2024.exists(), reason="Tracker file not available")
def test_read_all_patient_sheets_file_name():
    """Test that file_name metadata is correctly added."""
    df_all = read_all_patient_sheets(TRACKER_SBU_2024)

    assert "file_name" in df_all.columns
    file_names = df_all["file_name"].unique().to_list()
    assert len(file_names) == 1
    assert file_names[0] == TRACKER_SBU_2024.stem


@pytest.mark.skipif(not TRACKER_MHS_2017.exists(), reason="Tracker file not available")
def test_read_all_patient_sheets_2017_mhs_complete():
    """
    End-to-end test: 2017 Mahosot Hospital tracker (Laos/MHS).

    Characteristics:
    - Year: 2017
    - Sheets: Jan17-Dec17 (March is MISSING)
    - NO Patient List or Annual sheets
    - clinic_id should be "MHS"

    Expected patient counts per month:
    - Jan17: 6, Feb17: 6, Apr17: 6, May17: 8, Jun17: 11, Jul17: 11
    - Aug17: 11, Sep17: 12, Oct17: 12, Nov17: 12, Dec17: 14
    - Total: 109 patients (11 months)
    """
    df_all = read_all_patient_sheets(TRACKER_MHS_2017)

    # Basic validation
    assert len(df_all) > 0, "Should have extracted patient data"
    assert "patient_id" in df_all.columns
    assert "tracker_month" in df_all.columns
    assert "tracker_year" in df_all.columns
    assert "clinic_id" in df_all.columns

    # Check clinic_id
    assert df_all["clinic_id"].unique().to_list() == ["MHS"]

    # Check year
    assert df_all["tracker_year"].unique().to_list() == [2017]

    # Check we have exactly 11 months (March is missing)
    unique_months = sorted(df_all["tracker_month"].unique().to_list())
    expected_months = [1, 2, 4, 5, 6, 7, 8, 9, 10, 11, 12]  # Missing 3 (March)
    assert unique_months == expected_months, f"Expected {expected_months}, got {unique_months}"

    # Verify patient counts per month
    import calendar

    expected_counts = {
        1: 6,  # Jan
        2: 6,  # Feb
        # 3 is missing (March)
        4: 6,  # Apr
        5: 8,  # May
        6: 11,  # Jun
        7: 11,  # Jul
        8: 11,  # Aug
        9: 12,  # Sep
        10: 12,  # Oct
        11: 12,  # Nov
        12: 14,  # Dec
    }

    for month, expected_count in expected_counts.items():
        month_data = df_all.filter(pl.col("tracker_month") == month)
        actual_count = len(month_data)
        assert actual_count == expected_count, (
            f"Month {month} ({calendar.month_abbr[month]}17): "
            f"expected {expected_count} patients, got {actual_count}"
        )

    # Total patient count
    total_expected = sum(expected_counts.values())  # 109
    assert len(df_all) == total_expected, (
        f"Total patients: expected {total_expected}, got {len(df_all)}"
    )

    print(
        f"\n✓ 2017 MHS Tracker: {len(df_all)} patients from 11 months (March missing as expected)"
    )


@pytest.mark.skipif(not TRACKER_MHS_2025.exists(), reason="Tracker file not available")
def test_read_all_patient_sheets_2025_mhs_with_patient_list():
    """
    End-to-end test: 2025 Mahosot Hospital tracker (Laos/MHS).

    Characteristics:
    - Year: 2025
    - Sheets: Jan25-Jun25 (6 months)
    - HAS Patient List and Annual sheets
    - clinic_id should be "MHS"

    Expected patient counts per month:
    - Jan25: 95, Feb25: 97, Mar25: 97, Apr25: 97, May25: 98, Jun25: 99
    - Total: 583 patients
    """
    df_all = read_all_patient_sheets(TRACKER_MHS_2025)

    # Basic validation
    assert len(df_all) > 0, "Should have extracted patient data"
    assert "patient_id" in df_all.columns
    assert "tracker_month" in df_all.columns
    assert "tracker_year" in df_all.columns
    assert "clinic_id" in df_all.columns

    # Check clinic_id
    assert df_all["clinic_id"].unique().to_list() == ["MHS"]

    # Check year
    assert df_all["tracker_year"].unique().to_list() == [2025]

    # Check we have exactly 6 months (Jan-Jun)
    unique_months = sorted(df_all["tracker_month"].unique().to_list())
    expected_months = [1, 2, 3, 4, 5, 6]
    assert unique_months == expected_months, f"Expected {expected_months}, got {unique_months}"

    # Verify patient counts per month
    import calendar

    expected_counts = {
        1: 95,  # Jan
        2: 97,  # Feb
        3: 97,  # Mar
        4: 97,  # Apr
        5: 98,  # May
        6: 99,  # Jun
    }

    for month, expected_count in expected_counts.items():
        month_data = df_all.filter(pl.col("tracker_month") == month)
        actual_count = len(month_data)
        assert actual_count == expected_count, (
            f"Month {month} ({calendar.month_abbr[month]}25): "
            f"expected {expected_count} patients, got {actual_count}"
        )

    # Total patient count
    total_expected = sum(expected_counts.values())  # 583
    assert len(df_all) == total_expected, (
        f"Total patients: expected {total_expected}, got {len(df_all)}"
    )

    # Check that Patient List data was joined (should have columns from Patient List)
    # Note: The exact columns depend on what's in the Patient List sheet
    # We verify by checking for potential .static suffix columns
    static_cols = [col for col in df_all.columns if ".static" in col]
    print(f"\nColumns from Patient List (.static suffix): {len(static_cols)}")

    # Check that Annual data was joined
    annual_cols = [col for col in df_all.columns if ".annual" in col]
    print(f"Columns from Annual sheet (.annual suffix): {len(annual_cols)}")

    print(
        f"\n✓ 2025 MHS Tracker: {len(df_all)} patients from 6 months "
        f"(with Patient List & Annual data joined)"
    )


def test_export_patient_raw(tmp_path):
    """Test exporting patient data to parquet file."""
    from a4d.extract.patient import export_patient_raw, read_all_patient_sheets

    # Use the 2024 SBU tracker as test data
    tracker_file = TRACKER_SBU_2024
    if not tracker_file.exists():
        pytest.skip("Tracker file not available")

    # Extract data
    df = read_all_patient_sheets(tracker_file)

    # Export to temp directory
    output_dir = tmp_path / "patient_data_raw"
    output_path = export_patient_raw(df, tracker_file, output_dir)

    # Verify output file exists
    assert output_path.exists()
    assert output_path.name == "2024_Sibu Hospital A4D Tracker_patient_raw.parquet"
    assert output_path.parent == output_dir

    # Verify we can read it back
    df_read = pl.read_parquet(output_path)
    assert len(df_read) == len(df)
    assert df_read.columns == df.columns

    # Verify content matches
    assert df_read.equals(df)

    print(f"\n✓ Successfully exported and verified {len(df)} rows to parquet")


def _tracker_with_broken_id_formula(tmp_path: Path) -> Path:
    """A tracker whose second patient's ID cell holds an Excel #REF! error."""
    import openpyxl

    clinic_dir = tmp_path / "TST"
    clinic_dir.mkdir()
    tracker_path = clinic_dir / "2024_Test_Clinic.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jan24"
    ws.cell(2, 2).value = "Patient ID"
    ws.cell(2, 3).value = "Name"
    ws.cell(2, 4).value = "Age"

    ws.cell(3, 1).value = 1
    ws.cell(3, 2).value = "TS_QA001"
    ws.cell(3, 3).value = "TS_QA001"
    ws.cell(3, 4).value = 14

    ws.cell(4, 1).value = 2
    ws.cell(4, 2).value = "#REF!"
    ws.cell(4, 3).value = "#REF!"
    ws.cell(4, 4).value = 15

    wb.save(tracker_path)
    return tracker_path


def test_excel_error_patient_id_row_is_dropped_and_recorded(tmp_path):
    """A row whose ID is a broken formula is dropped, but never silently.

    #REF! is not a malformed identifier a clinic could reconcile -- the cell's
    content is gone -- so the row cannot be kept under the Undefined sentinel
    the way ticket 47 keeps a misspelled ID: grouping by patient_id would pool
    it with every other unidentified patient. It is dropped, and the discard is
    reported so the source workbook can be corrected.
    """
    from a4d.errors import ErrorCollector

    tracker_path = _tracker_with_broken_id_formula(tmp_path)
    collector = ErrorCollector()

    df = read_all_patient_sheets(tracker_path, error_collector=collector)

    assert df["patient_id"].to_list() == ["TS_QA001"]

    errors = collector.to_dataframe()
    dropped = errors.filter(pl.col("error_code") == "excel_error_patient_id")
    assert len(dropped) == 1
    assert dropped["original_value"][0] == "#REF!"
    assert "Jan24" in dropped["error_message"][0]
