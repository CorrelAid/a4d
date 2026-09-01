"""Build the golden-master tracker: one synthetic workbook shaped like a real one.

The workbook carries the 2026 template's real structure -- the sheets, the
two-row headers, the title block above them, the product blocks stacked above
the patient block on a month sheet -- because that structure is what the
extractor navigates, and a fixture that flattens it exercises none of that
navigation.

Everything a person could be identified by is invented: the patients, their
IDs, their names, every reading and date. What is *not* invented is anything
drawn from a controlled vocabulary -- province, patient status, support level,
insulin regimen, product names -- since those are matched against
``reference_data/`` and a made-up value would exercise the failure path
instead of the normal one.

The workbook is built here rather than committed as a binary so that a change
to the fixture shows up as a readable diff.
"""

from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# A real clinic folder, so clinic_id resolves against reference_data as it
# does in production. The patients inside it are not this clinic's.
CLINIC_FOLDER = "KBH"
TRACKER_NAME = "2026_Golden Master A4D Tracker.xlsx"
MONTH_SHEETS = ("Jan26", "Feb26")

_TITLE_BLOCK = [
    (2, 4, "CLINIC SUPPORT PROGRAM"),
    (3, 4, "GOLDEN MASTER HOSPITAL"),
    (4, 4, "CAMBODIA"),
]

# (patient_id, name, province, sex, dob, dx_date, dx_age, dka, hba1c, fbg,
#  recruitment, consent, phone, insurance, regimen, bgm, insulin)
_PATIENTS = [
    (
        "KH_QB001",
        "Golden One",
        "Phnom Penh",
        "F",
        date(2008, 3, 14),
        date(2017, 6, 2),
        9,
        "Y",
        7.4,
        149.9,
        date(2018, 7, 1),
        "Y",
        "012 000 001",
        "Yes",
        "Basal-Bolus (MDI)",
        "Y",
        "Y",
    ),
    (
        "KH_QB002",
        "Golden Two",
        "Kandal",
        "M",
        date(2011, 11, 2),
        date(2019, 1, 20),
        7,
        "N",
        9.1,
        210.0,
        date(2019, 3, 1),
        "Y",
        "012 000 002",
        "No",
        "Premixed 30/70 BD",
        "Y",
        "N",
    ),
    (
        "KH_QB003",
        "Golden Three",
        "Kampong Cham",
        "F",
        date(2006, 7, 30),
        date(2015, 9, 9),
        9,
        "Y",
        11.2,
        305.5,
        date(2016, 2, 1),
        "N",
        "012 000 003",
        "Yes",
        "Self-mixed BD",
        "N",
        "Y",
    ),
    (
        "KH_QB004",
        "Golden Four",
        "Phnom Penh",
        "M",
        date(2013, 1, 5),
        date(2021, 4, 18),
        8,
        "N",
        8.0,
        178.2,
        date(2021, 6, 1),
        "Y",
        "012 000 004",
        "Yes",
        "Modified conventional TID",
        "Y",
        "Y",
    ),
]


def _write_title_block(ws: Worksheet) -> None:
    for row, col, value in _TITLE_BLOCK:
        ws.cell(row, col).value = value


def _write_row(ws: Worksheet, row: int, values: dict[int, object]) -> None:
    for col, value in values.items():
        ws.cell(row, col).value = value


def _build_patient_list(wb: openpyxl.Workbook) -> None:
    """The roster sheet: one row per patient, demographics that never change."""
    ws = wb.create_sheet("Patient List")
    _write_title_block(ws)
    ws.cell(6, 2).value = "PATIENT RECRUITMENT"
    _write_row(
        ws,
        8,
        {
            2: "Patient \nID",
            3: "Patient Name",
            4: "Province",
            5: "Gender",
            6: "D.O.B.",
            7: "Date of T1D\nDiagnosis",
            8: "Age at\nDiagnosis*",
            9: "DKA @\nDiagnosis",
            10: "Baseline HbA1c",
            11: "Baseline \nFBG",
            12: "Date of Recruitment ",
            13: "Lost Patients \nSummary",
            15: "Patient \nConsent \nCollected",
            17: "Phone Number",
            18: "Insurance Card Status",
            19: "Current Insulin Regimen",
            20: "BGM \nA4D",
            21: "Insulin\nA4D",
        },
    )
    _write_row(
        ws,
        9,
        {
            6: "(dd-mmm-yyyy)",
            7: "(dd-mmm-yyyy)",
            9: "Y OR N",
            10: "%",
            11: "mg/dL",
            12: "(mmm-yyyy)",
            13: "Lost Date\n(dd-mmm-yyyy)",
            14: "Status OUT",
        },
    )
    for offset, p in enumerate(_PATIENTS):
        _write_row(
            ws,
            10 + offset,
            {
                1: offset + 1,
                2: p[0],
                3: p[1],
                4: p[2],
                5: p[3],
                6: p[4],
                7: p[5],
                8: p[6],
                9: p[7],
                10: p[8],
                11: p[9],
                12: p[10],
                15: p[11],
                17: p[12],
                18: p[13],
                19: p[14],
                20: p[15],
                21: p[16],
            },
        )


def _build_annual(wb: openpyxl.Workbook) -> None:
    """The once-a-year sheet: complication screening under a merged header.

    The merged header spanning the screening block is deliberate -- it is the
    shape that once caused screening results to be dropped entirely.
    """
    ws = wb.create_sheet("Annual")
    _write_title_block(ws)
    ws.cell(6, 2).value = "ANNUAL PATIENT UPDATE"
    ws.cell(8, 7).value = "Complication Screening"
    ws.merge_cells(start_row=8, start_column=7, end_row=8, end_column=23)
    _write_row(
        ws,
        9,
        {
            2: "Patient \nID*",
            3: "Patient Name*",
            4: "Patient Status",
            5: "Level of Education\nOr Occupation",
            7: "Blood Pressure",
            10: "Kidney Function Test",
            12: "Eye Exam",
            14: "Foot Exam",
            16: "Lipid Profile",
            21: "Thyroid Function Test",
            24: "Remarks",
            25: "DM Complication (Y/ N)",
            29: "Family History",
            30: "Other health or psychosocial issues",
        },
    )
    _write_row(
        ws,
        10,
        {
            6: "Date Updated\n(dd-mmm-yyyy)",
            7: "Date\n(dd-mmm-yyyy)",
            8: "Systolic \n(mmHg)",
            9: "Diastolic (mmHg)",
            10: "Date\n(dd-mmm-yyyy)",
            11: "UACR\n(mg/g)",
            12: "Date\n(dd-mmm-yyyy)",
            13: "Result",
            14: "Date\n(dd-mmm-yyyy)",
            15: "Result",
            16: "Date\n(dd-mmm-yyyy)",
            17: "Triglycerides",
            18: "Cholesterol",
            19: "LDL mg/dL",
            20: "HDL mg/dL",
            21: "Date\n(dd-mmm-yyyy)",
            22: "FT4 pmol/L",
            23: "TSH\n(mlU/L)",
            25: "Eye",
            26: "Kidney",
            27: "Others",
            28: "Remarks",
        },
    )
    screening = [
        (
            date(2026, 2, 10),
            110,
            70,
            date(2026, 2, 10),
            12.0,
            date(2026, 2, 11),
            "Normal",
            date(2026, 2, 11),
            "Normal",
            date(2026, 2, 12),
            90,
            150,
            80,
            55,
            date(2026, 2, 12),
            14.0,
            2.1,
            "N",
            "N",
            "N",
        ),
        (
            date(2026, 3, 3),
            125,
            82,
            date(2026, 3, 3),
            35.0,
            date(2026, 3, 4),
            "Abnormal",
            date(2026, 3, 4),
            "Normal",
            date(2026, 3, 5),
            180,
            210,
            130,
            40,
            date(2026, 3, 5),
            12.5,
            4.4,
            "Y",
            "N",
            "N",
        ),
    ]
    for offset, (p, s) in enumerate(zip(_PATIENTS[:2], screening, strict=True)):
        _write_row(
            ws,
            11 + offset,
            {
                1: offset + 1,
                2: p[0],
                3: p[1],
                4: "Active",
                5: "Secondary school",
                6: date(2026, 1, 15),
                7: s[0],
                8: s[1],
                9: s[2],
                10: s[3],
                11: s[4],
                12: s[5],
                13: s[6],
                14: s[7],
                15: s[8],
                16: s[9],
                17: s[10],
                18: s[11],
                19: s[12],
                20: s[13],
                21: s[14],
                22: s[15],
                23: s[16],
                24: "Annual review completed",
                25: s[17],
                26: s[18],
                27: s[19],
                28: "-",
                29: "Father T2D",
                30: "None reported",
            },
        )


_PRODUCT_HEADER = {
    2: "Product",
    5: "Entry Date",
    6: "Balance",
    7: " Units Received",
    8: "Received From",
    10: "Units Released",
    11: "Released To (Drop Down)",
    13: "REMARKS",
}

# Two stock blocks per month sheet, as real trackers stack them: the product
# is named once on the block's first row and the movements follow beneath it
# with the name left blank, which is the shape the extractor forward-fills.
_STOCK = {
    "Jan26": [
        (
            "CareSens N Strips (50s)",
            [
                (date(2026, 1, 1), 60, None, None, None, None, "Opening balance"),
                (date(2026, 1, 8), 100, 40, "A4D HQ", None, None, None),
                (date(2026, 1, 15), 98, None, None, 2, "KH_QB001", None),
                (date(2026, 1, 22), 96, None, None, 2, "KH_QB002", None),
            ],
        ),
        (
            "Accu-Chek Softclix Lancet (200s)",
            [
                (date(2026, 1, 1), 12, None, None, None, None, None),
                (date(2026, 1, 20), 11, None, None, 1, "KH_QB003", None),
            ],
        ),
    ],
    "Feb26": [
        (
            "CareSens N Strips (50s)",
            [
                (date(2026, 2, 1), 96, None, None, None, None, None),
                (date(2026, 2, 12), 94, None, None, 2, "KH_QB004", None),
            ],
        ),
        (
            "Accu-Chek Softclix Lancet (200s)",
            [
                (date(2026, 2, 1), 11, None, None, None, None, None),
                (date(2026, 2, 18), 31, 20, "A4D HQ", None, None, "Quarterly delivery"),
            ],
        ),
    ],
}

_MONTH_PATIENT_HEADER_1 = {
    2: "Patient \nID*",
    3: "Patient Name*",
    4: "Clinic Visit",
    5: "Last Clinic \nVisit",
    6: "Remote Follow Up",
    7: "Last Remote \nFollow Up",
    8: "Patient Status",
    9: "Level of Support",
    10: "Age*",
    11: "Baseline \nHbA1c*",
    12: "Updated \nHbA1c",
    14: "Updated FBG",
    16: "Body \nWeight",
    17: "Height",
    18: "BMI*",
    19: "Date of \nBMI",
    20: "Current Month Hospitalisation \n(DKA, Hyper, Hypo, Other)",
    22: "Current Patient\nObservations",
    25: "Insulin Regimen",
    26: "Testing Frequency",
    27: "TOTAL Insulin Units ",
    28: "Number of insulin injections",
    29: "Pre-mixed",
    30: "Short-acting",
    31: "Intermediate-acting",
    32: "Rapid-acting",
    33: "Long-acting",
}

_MONTH_PATIENT_HEADER_2 = {
    5: "(dd-mmm-yyyy)",
    7: "(dd-mmm-yyyy)",
    9: "Drop Down",
    10: "On Reporting",
    11: "%",
    12: "%",
    13: "(dd-mmm-yyyy)",
    14: "mg/dL",
    15: "(dd-mmm-yyyy)",
    16: "kg",
    17: "meters",
    19: "(dd-mmm-yyyy)",
    20: "Drop Down",
    21: "(dd-mmm-yyyy)",
    23: "Category",
    26: "per day",
    27: "per day",
    28: "per day",
    29: "Select",
    30: "Select",
    31: "Select",
    32: "Select",
    33: "Select",
}

# One monthly reading set per patient per month. Values are invented; the
# fields drawn from a vocabulary (status, support level, hospitalisation
# cause, observation category) are real ones.
# One reading set per patient per month, keyed by field so the layout below
# cannot silently shift a value into the wrong column. Numbers are invented;
# the fields backed by a dropdown (status, support level, hospitalisation
# cause, observation category, insulin regimen and the five insulin subtype
# flags) carry values the reference data actually allows.
_MONTHLY: dict[str, list[dict[str, object]]] = {
    "Jan26": [
        {
            "patient_id": "KH_QB001",
            "visit": "Y",
            "visit_date": date(2026, 1, 12),
            "remote": "N",
            "status": "Active",
            "support": "Standard",
            "age": 17,
            "hba1c_baseline": 7.4,
            "hba1c_updated": 7.8,
            "hba1c_date": date(2026, 1, 12),
            "fbg": 130.0,
            "weight": 52.0,
            "height": 1.62,
            "bmi": 19.8,
            "bmi_date": date(2026, 1, 12),
            "regimen": "Basal-Bolus (MDI)",
            "testing_frequency": 4,
            "insulin_units": 32,
            "injections": 4,
            "rapid": "Y",
            "long": "Y",
        },
        {
            "patient_id": "KH_QB002",
            "visit": "N",
            "remote": "Y",
            "remote_date": date(2026, 1, 9),
            "status": "Active Remote",
            "support": "Partial",
            "age": 14,
            "hba1c_baseline": 9.1,
            "hba1c_updated": 9.6,
            "hba1c_date": date(2026, 1, 9),
            "fbg": 190.0,
            "weight": 41.0,
            "height": 1.50,
            "bmi": 18.2,
            "bmi_date": date(2026, 1, 9),
            "hospitalisation": "HYPO",
            "observations": "Missed doses",
            "observations_category": "Clinic Follow Up",
            "regimen": "Premixed 30/70 BD",
            "testing_frequency": 2,
            "insulin_units": 28,
            "injections": 2,
            "premixed": "Y",
        },
        {
            "patient_id": "KH_QB003",
            "visit": "Y",
            "visit_date": date(2026, 1, 20),
            "remote": "N",
            "status": "Active",
            "support": "SAC",
            "age": 19,
            "hba1c_baseline": 11.2,
            "hba1c_updated": 10.4,
            "hba1c_date": date(2026, 1, 20),
            "fbg": 240.0,
            "weight": 55.0,
            "height": 1.58,
            "bmi": 22.0,
            "bmi_date": date(2026, 1, 20),
            "regimen": "Self-mixed BD",
            "testing_frequency": 3,
            "insulin_units": 40,
            "injections": 2,
            "short": "Y",
            "intermediate": "Y",
        },
        {
            "patient_id": "KH_QB004",
            "visit": "Y",
            "visit_date": date(2026, 1, 27),
            "remote": "N",
            "status": "Active Monitoring",
            "support": "Monitoring",
            "age": 13,
            "hba1c_baseline": 8.0,
            "hba1c_updated": 8.2,
            "hba1c_date": date(2026, 1, 27),
            "fbg": 165.0,
            "weight": 38.0,
            "height": 1.44,
            "bmi": 18.3,
            "bmi_date": date(2026, 1, 27),
            "regimen": "Modified conventional TID",
            "testing_frequency": 3,
            "insulin_units": 24,
            "injections": 3,
            "short": "Y",
            "rapid": "Y",
        },
    ],
    "Feb26": [
        {
            "patient_id": "KH_QB001",
            "visit": "Y",
            "visit_date": date(2026, 2, 9),
            "remote": "N",
            "status": "Active",
            "support": "Standard",
            "age": 17,
            "hba1c_baseline": 7.4,
            "hba1c_updated": 7.5,
            "hba1c_date": date(2026, 2, 9),
            "fbg": 124.0,
            "weight": 52.5,
            "height": 1.62,
            "bmi": 20.0,
            "bmi_date": date(2026, 2, 9),
            "regimen": "Basal-Bolus (MDI)",
            "testing_frequency": 4,
            "insulin_units": 33,
            "injections": 4,
            "rapid": "Y",
            "long": "Y",
        },
        {
            "patient_id": "KH_QB002",
            "visit": "Y",
            "visit_date": date(2026, 2, 14),
            "remote": "N",
            "status": "Active",
            "support": "Partial",
            "age": 14,
            "hba1c_baseline": 9.1,
            "hba1c_updated": 9.0,
            "hba1c_date": date(2026, 2, 14),
            "fbg": 180.0,
            "weight": 41.5,
            "height": 1.50,
            "bmi": 18.4,
            "bmi_date": date(2026, 2, 14),
            "regimen": "Premixed 30/70 BD",
            "testing_frequency": 2,
            "insulin_units": 28,
            "injections": 2,
            "premixed": "Y",
        },
        # Leaves the programme mid-year: every reading blank, status carried
        # by the dropdown alone. The shape a lost patient actually has.
        {
            "patient_id": "KH_QB003",
            "visit": "N",
            "remote": "N",
            "status": "Lost Follow Up",
            "support": "SAC",
            "age": 19,
            "hba1c_baseline": 11.2,
            "observations_category": "Status OUT",
        },
        {
            "patient_id": "KH_QB004",
            "visit": "Y",
            "visit_date": date(2026, 2, 24),
            "remote": "N",
            "status": "Active Monitoring",
            "support": "Monitoring",
            "age": 13,
            "hba1c_baseline": 8.0,
            "hba1c_updated": 7.9,
            "hba1c_date": date(2026, 2, 24),
            "fbg": 158.0,
            "weight": 38.5,
            "height": 1.45,
            "bmi": 18.3,
            "bmi_date": date(2026, 2, 24),
            "hospitalisation": "DKA",
            "observations": "Admitted 2 days",
            "observations_category": "Hospitalisation",
            "regimen": "Modified conventional TID",
            "testing_frequency": 3,
            "insulin_units": 25,
            "injections": 3,
            "short": "Y",
            "rapid": "Y",
        },
    ],
}

# Field -> the column it occupies in the month sheet's patient block.
_MONTH_PATIENT_COLUMNS = {
    "patient_id": 2,
    "name": 3,
    "visit": 4,
    "visit_date": 5,
    "remote": 6,
    "remote_date": 7,
    "status": 8,
    "support": 9,
    "age": 10,
    "hba1c_baseline": 11,
    "hba1c_updated": 12,
    "hba1c_date": 13,
    "fbg": 14,
    "weight": 16,
    "height": 17,
    "bmi": 18,
    "bmi_date": 19,
    "hospitalisation": 20,
    "observations": 22,
    "observations_category": 23,
    "regimen": 25,
    "testing_frequency": 26,
    "insulin_units": 27,
    "injections": 28,
    "premixed": 29,
    "short": 30,
    "intermediate": 31,
    "rapid": 32,
    "long": 33,
}

# The five subtype flags are a dropdown of Y or "-", never a number.
_SUBTYPE_FLAGS = ("premixed", "short", "intermediate", "rapid", "long")


def _build_month_sheet(wb: openpyxl.Workbook, sheet: str, month_end: date) -> None:
    """A month sheet: summary block, stacked stock blocks, then the patients."""
    ws = wb.create_sheet(sheet)
    _write_title_block(ws)
    _write_row(ws, 6, {2: "REPORT START DATE", 4: month_end.replace(day=1)})
    _write_row(ws, 7, {2: "REPORT END DATE", 4: month_end})
    _write_row(ws, 10, {2: "PROVINCE", 4: "Phnom Penh"})
    _write_row(ws, 11, {2: "CSP FOCAL PERSON", 4: "Golden Nurse"})
    ws.cell(14, 2).value = "MEDICAL SUPPLIES DISTRIBUTION"

    row = 16
    for product, movements in _STOCK[sheet]:
        _write_row(ws, row, _PRODUCT_HEADER)
        row += 1
        for index, (entry, balance, received, source, released, target, remark) in enumerate(
            movements
        ):
            _write_row(
                ws,
                row,
                {
                    2: product if index == 0 else None,
                    5: entry,
                    6: balance,
                    7: received,
                    8: source,
                    10: released,
                    11: target,
                    13: remark,
                },
            )
            row += 1
        row += 2  # blank rows between stock blocks, as in the real template

    # The marker the extractor uses to find where the stock section ends and
    # the patient block begins; without it the whole sheet's stock is skipped.
    ws.cell(row + 1, 2).value = "PATIENT DATA SUMMARY"

    header_row = row + 3
    _write_row(ws, header_row, _MONTH_PATIENT_HEADER_1)
    _write_row(ws, header_row + 1, _MONTH_PATIENT_HEADER_2)
    for offset, reading in enumerate(_MONTHLY[sheet]):
        row_values = {_MONTH_PATIENT_COLUMNS[field]: value for field, value in reading.items()}
        row_values[1] = offset + 1
        row_values[_MONTH_PATIENT_COLUMNS["name"]] = f"Golden {reading['patient_id']}"
        for flag in _SUBTYPE_FLAGS:
            row_values.setdefault(_MONTH_PATIENT_COLUMNS[flag], "-")
        _write_row(ws, header_row + 2 + offset, row_values)


def build_golden_tracker(root: Path) -> Path:
    """Write the golden-master workbook under ``root`` and return its path."""
    clinic_dir = root / CLINIC_FOLDER
    clinic_dir.mkdir(parents=True, exist_ok=True)
    tracker_path = clinic_dir / TRACKER_NAME

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _build_patient_list(wb)
    _build_annual(wb)
    for sheet, month_end in zip(MONTH_SHEETS, (date(2026, 1, 31), date(2026, 2, 28)), strict=True):
        _build_month_sheet(wb, sheet, month_end)
    wb.save(tracker_path)
    return tracker_path
