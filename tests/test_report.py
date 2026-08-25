"""Tests for the findings report workbook.

The workbook is the deliverable A4D staff act on, so these pin the three
properties that make it actionable: the summary ranks by the category that
means "a human must fix this", the glossary covers every code whether or not
this run produced one, and one tracker can be drilled into on its own.
"""

from datetime import datetime

import polars as pl
import pytest
import xlsxwriter
from openpyxl import load_workbook

from a4d.findings import FINDING_GLOSSARY
from a4d.report import (
    build_findings_report,
    glossary_frame,
    load_findings,
    summarise_by_tracker,
)


def _findings(rows: list[dict]) -> pl.DataFrame:
    defaults = {
        "file_name": "2024_Clinic_A",
        "arm": "patient",
        "sheet_name": "Jan24",
        "patient_id": "KH_KB001",
        "column": "age",
        "original_value": "abc",
        "message": "could not convert",
        "error_code": "type_conversion",
        "category": "data_lost",
        "stage": "clean",
        "function_name": "safe_convert_column",
        "tracker_year": 2024,
        "tracker_month": 1,
        "timestamp": datetime(2026, 8, 25, 12, 0, 0),
    }
    return pl.DataFrame([{**defaults, **row} for row in rows])


class TestSummary:
    def test_the_tracker_needing_a_human_is_first(self):
        """A staff member opens this to find which workbooks to correct, so a
        tracker with one workbook defect outranks one with a thousand
        recoveries."""
        findings = _findings(
            [
                *[{"file_name": "noisy", "category": "recovered"}] * 50,
                {"file_name": "broken", "category": "fix_workbook"},
            ]
        )

        summary = summarise_by_tracker(findings)

        assert summary["file_name"].to_list() == ["broken", "noisy"]

    def test_a_tracker_counts_each_category_separately(self):
        findings = _findings(
            [
                {"category": "fix_workbook"},
                {"category": "fix_workbook"},
                {"category": "data_lost"},
                {"category": "recovered"},
            ]
        )

        row = summarise_by_tracker(findings).row(0, named=True)

        assert (row["fix_workbook"], row["data_lost"], row["recovered"]) == (2, 1, 1)
        assert row["total"] == 4

    def test_a_tracker_found_by_both_arms_names_both(self):
        """The summary is per workbook, not per workbook-and-arm: a clinic
        correcting a tracker fixes one file."""
        findings = _findings([{"arm": "patient"}, {"arm": "product"}])

        assert summarise_by_tracker(findings).row(0, named=True)["arms"] == "patient, product"

    def test_an_empty_run_still_has_the_summary_columns(self):
        """A clean run must not produce a workbook with a different shape."""
        summary = summarise_by_tracker(_findings([]).clear())

        assert "fix_workbook" in summary.columns
        assert summary.is_empty()


class TestGlossary:
    def test_every_code_appears_even_when_this_run_produced_none(self):
        """Someone looking up a code they saw last month should find it."""
        glossary = glossary_frame(_findings([{"error_code": "type_conversion"}]))

        assert set(glossary["error_code"].to_list()) == set(FINDING_GLOSSARY)

    def test_a_code_that_did_not_fire_reads_zero_rather_than_blank(self):
        glossary = glossary_frame(_findings([{"error_code": "type_conversion"}]))
        counts = dict(zip(glossary["error_code"], glossary["count_in_this_run"], strict=True))

        assert counts["type_conversion"] == 1
        assert counts["balance_reconciliation"] == 0

    def test_workbook_defects_are_listed_before_recoveries(self):
        categories = glossary_frame(_findings([]).clear())["category"].to_list()

        assert categories.index("fix_workbook") < categories.index("recovered")


class TestWorkbook:
    def test_it_writes_the_three_agreed_sheets(self, tmp_path):
        out = build_findings_report(_findings([{}]), tmp_path / "findings.xlsx")

        assert [sheet.title for sheet in _sheets(out)] == ["Summary", "Findings", "Glossary"]

    def test_the_findings_sheet_leads_with_actionability(self, tmp_path):
        """Category first so a reader never scrolls right to learn whether a
        row matters."""
        out = build_findings_report(_findings([{}]), tmp_path / "findings.xlsx")

        assert _header(out, "Findings")[:3] == ["category", "file_name", "sheet_name"]

    def test_drilling_into_one_tracker_excludes_the_others(self, tmp_path):
        findings = _findings([{"file_name": "2024_Wanted"}, {"file_name": "2024_Other"}])

        out = build_findings_report(findings, tmp_path / "one.xlsx", tracker="Wanted")

        assert _column(out, "Findings", "file_name") == ["2024_Wanted"]

    def test_the_glossary_still_covers_every_code_when_drilled_down(self, tmp_path):
        """Filtering to one tracker must not shrink the reference material."""
        findings = _findings([{"file_name": "2024_Wanted"}, {"file_name": "2024_Other"}])

        out = build_findings_report(findings, tmp_path / "one.xlsx", tracker="Wanted")

        assert len(_column(out, "Glossary", "error_code")) == len(FINDING_GLOSSARY)

    def test_an_unknown_tracker_is_an_error_not_an_empty_workbook(self, tmp_path):
        """An empty file looks like a clean tracker; a typo should not."""
        with pytest.raises(ValueError, match="No findings"):
            build_findings_report(_findings([{}]), tmp_path / "none.xlsx", tracker="nonexistent")

    def test_a_missing_findings_table_says_how_to_produce_one(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="a4d run"):
            load_findings(tmp_path / "absent.parquet")


def _sheets(path):
    return load_workbook(path).worksheets


def _header(path, sheet_name: str) -> list[str]:
    sheet = load_workbook(path)[sheet_name]
    return [cell.value for cell in next(sheet.iter_rows(max_row=1))]


def _column(path, sheet_name: str, column: str) -> list:
    sheet = load_workbook(path)[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    index = list(rows[0]).index(column)
    return [row[index] for row in rows[1:] if row[index] is not None]


def test_xlsxwriter_is_available():
    """The workbook is written with xlsxwriter, added for this report."""
    assert xlsxwriter.__version__


def test_the_summary_carries_no_column_the_pipeline_never_fills(tmp_path):
    """`tracker_year` is declared on the findings table but never populated
    (0 of 122,590 on the real run), so shipping it would be a blank column in
    the deliverable."""
    out = build_findings_report(_findings([{}]), tmp_path / "findings.xlsx")

    assert "tracker_year" not in _header(out, "Summary")
