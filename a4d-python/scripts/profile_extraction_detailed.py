"""Detailed timing breakdown of extraction phases."""

import time
from pathlib import Path

from openpyxl import load_workbook

TRACKER_2024 = Path(
    "/Volumes/USB SanDisk 3.2Gen1 Media/A4D/data/a4dphase2_upload/"
    "Malaysia/SBU/2024_Sibu Hospital A4D Tracker.xlsx"
)
TRACKER_2019 = Path(
    "/Volumes/USB SanDisk 3.2Gen1 Media/A4D/data/a4dphase2_upload/"
    "Malaysia/PNG/2019_Penang General Hospital A4D Tracker_DC.xlsx"
)


def profile_extraction_phases(tracker_file, sheet_name, year):
    """Profile each phase of extraction separately.

    NOTE: This is the OPTIMIZED single-pass version that matches the current implementation.
    """
    print(f"\n{'=' * 80}")
    print(f"Profiling: {tracker_file.name} - {sheet_name}")
    print("=" * 80)

    timings = {}

    # Phase 1: Load workbook (read-only for optimal performance)
    t0 = time.perf_counter()
    wb = load_workbook(
        tracker_file,
        read_only=True,
        data_only=True,
        keep_vba=False,
        keep_links=False,
    )
    ws = wb[sheet_name]
    t1 = time.perf_counter()
    timings["1. Load workbook (read-only)"] = t1 - t0

    # Phase 2: Find data start row
    t0 = time.perf_counter()
    data_start_row = None
    for row_idx, (cell_value,) in enumerate(
        ws.iter_rows(min_col=1, max_col=1, values_only=True), start=1
    ):
        if cell_value is not None:
            data_start_row = row_idx
            break
    t1 = time.perf_counter()
    timings["2. Find data start row"] = t1 - t0

    # Phase 3: Read headers
    t0 = time.perf_counter()
    header_row_1 = data_start_row - 1
    header_row_2 = data_start_row - 2

    max_cols = 100
    header_1_raw = list(
        ws.iter_rows(
            min_row=header_row_1,
            max_row=header_row_1,
            min_col=1,
            max_col=max_cols,
            values_only=True,
        )
    )[0]
    header_2_raw = list(
        ws.iter_rows(
            min_row=header_row_2,
            max_row=header_row_2,
            min_col=1,
            max_col=max_cols,
            values_only=True,
        )
    )[0]

    # Trim to actual width
    last_col = max_cols
    for i in range(len(header_1_raw) - 1, -1, -1):
        if header_1_raw[i] is not None or header_2_raw[i] is not None:
            last_col = i + 1
            break

    header_1 = list(header_1_raw[:last_col])
    header_2 = list(header_2_raw[:last_col])
    t1 = time.perf_counter()
    timings["3. Read headers"] = t1 - t0

    # Phase 4: Merge headers with forward-fill logic
    t0 = time.perf_counter()
    import re

    headers = []
    prev_h2 = None  # Track previous h2 for horizontal merges

    for h1, h2 in zip(header_1, header_2, strict=True):
        if h1 and h2:
            headers.append(f"{h2} {h1}".strip())
            prev_h2 = h2
        elif h2:
            headers.append(str(h2).strip())
            prev_h2 = h2
        elif h1:
            if prev_h2:
                # Horizontally merged cell: fill forward
                headers.append(f"{prev_h2} {h1}".strip())
            else:
                headers.append(str(h1).strip())
        else:
            headers.append(None)
            prev_h2 = None

    headers = [re.sub(r"\s+", " ", h.replace("\n", " ")) if h else None for h in headers]
    t1 = time.perf_counter()
    timings["4. Merge headers"] = t1 - t0

    # Phase 5: Read data rows
    t0 = time.perf_counter()
    data = []
    for row in ws.iter_rows(
        min_row=data_start_row,
        max_row=ws.max_row,
        min_col=1,
        max_col=len(headers),
        values_only=True,
    ):
        if all(cell is None for cell in row):
            break
        if row[0] is None:
            continue
        data.append(row)
    t1 = time.perf_counter()
    timings["5. Read data rows"] = t1 - t0

    # Phase 6: Close workbook
    t0 = time.perf_counter()
    wb.close()
    t1 = time.perf_counter()
    timings["6. Close workbook"] = t1 - t0

    # Phase 7: Build DataFrame
    t0 = time.perf_counter()
    import polars as pl

    valid_cols = [(i, h) for i, h in enumerate(headers) if h]
    valid_indices = [i for i, _ in valid_cols]
    valid_headers = [h for _, h in valid_cols]
    filtered_data = [[row[i] for i in valid_indices] for row in data]

    df = pl.DataFrame(
        {
            header: [str(row[i]) if row[i] is not None else None for row in filtered_data]
            for i, header in enumerate(valid_headers)
        }
    )
    t1 = time.perf_counter()
    timings["7. Build Polars DataFrame"] = t1 - t0

    # Print results
    total_time = sum(timings.values())
    print(f"\nExtracted: {len(df)} rows × {len(df.columns)} columns")
    print(f"Total time: {total_time:.3f}s\n")
    print(f"{'Phase':<40} {'Time (s)':<12} {'% of Total':<12}")
    print("-" * 64)

    for phase, duration in timings.items():
        pct = (duration / total_time) * 100
        print(f"{phase:<40} {duration:>10.3f}s  {pct:>10.1f}%")

    return timings, total_time


if __name__ == "__main__":
    # Test 2024 tracker
    timings_2024, total_2024 = profile_extraction_phases(TRACKER_2024, "Jan24", 2024)

    # Test 2019 tracker
    timings_2019, total_2019 = profile_extraction_phases(TRACKER_2019, "Feb19", 2019)

    print("\n" + "=" * 80)
    print("SUMMARY")
    print("=" * 80)
    print(f"2024 tracker total: {total_2024:.3f}s")
    print(f"2019 tracker total: {total_2019:.3f}s")
    print("\nSlowest phases across both trackers:")
    all_timings = {}
    for phase in timings_2024:
        all_timings[phase] = (timings_2024[phase] + timings_2019[phase]) / 2

    for phase, avg_time in sorted(all_timings.items(), key=lambda x: x[1], reverse=True)[:5]:
        print(f"  {phase:<40} avg: {avg_time:.3f}s")
