"""Patient data extraction from Excel tracker files.

This module handles reading patient data from Excel trackers, which have
evolved over the years with different formats and structures.
"""

import calendar
import re
import warnings
from pathlib import Path

import polars as pl
from loguru import logger
from openpyxl import load_workbook

from a4d.errors import ErrorCollector
from a4d.reference.synonyms import ColumnMapper, load_patient_mapper

# Suppress openpyxl warnings about unsupported Excel features
# We only read data, so these warnings are not actionable
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


def get_tracker_year(tracker_file: Path, month_sheets: list[str]) -> int:
    """Extract tracker year from month sheet names or filename.

    Tries to parse year from month sheet names (e.g., "Jan24" -> 2024).
    Falls back to extracting from filename if parsing fails.
    Validates year is in reasonable range (2017-2030).

    Args:
        tracker_file: Path to the tracker Excel file
        month_sheets: List of month sheet names

    Returns:
        Year of the tracker (e.g., 2024)

    Raises:
        ValueError: If year cannot be determined or is out of valid range

    Example:
        >>> get_tracker_year(Path("2024_Clinic.xlsx"), ["Jan24", "Feb24"])
        2024
    """
    for sheet in month_sheets:
        match = re.search(r"(\d{2})$", sheet)
        if match:
            year_suffix = int(match.group(1))
            year = 2000 + year_suffix  # Assume 20xx until 2100
            logger.debug(f"Parsed year {year} from sheet name '{sheet}'")

            if not (2017 <= year <= 2030):  # Match R pipeline validation
                raise ValueError(
                    f"Year {year} is out of valid range (2017-2030). "
                    f"Parsed from sheet name '{sheet}'"
                )

            return year

    match = re.search(r"(\d{4})", tracker_file.name)
    if match:
        year = int(match.group(1))
        logger.debug(f"Parsed year {year} from filename '{tracker_file.name}'")

        if not (2017 <= year <= 2030):  # Match R pipeline validation
            raise ValueError(
                f"Year {year} is out of valid range (2017-2030). "
                f"Parsed from filename '{tracker_file.name}'"
            )
        
        return year

    raise ValueError(
        f"Could not determine year from month sheets {month_sheets} or filename {tracker_file.name}"
    )


def find_month_sheets(workbook) -> list[str]:
    """Find all month sheets in the tracker workbook.

    Month sheets are identified by matching against month abbreviations
    (Jan, Feb, Mar, etc.) and sorted by month number for consistent processing.

    Args:
        workbook: openpyxl Workbook object

    Returns:
        List of month sheet names found in the workbook, sorted by month number
        (Jan=1, Feb=2, ..., Dec=12)

    Example:
        >>> wb = load_workbook("tracker.xlsx")
        >>> find_month_sheets(wb)
        ['Jan24', 'Feb24', 'Mar24', ...]
    """
    month_abbrs = list(calendar.month_abbr)[1:]  # ['Jan', 'Feb', ...]
    month_sheets = []

    for sheet_name in workbook.sheetnames:
        if any(sheet_name.startswith(abbr) for abbr in month_abbrs):
            month_sheets.append(sheet_name)

    def get_month_number(sheet_name: str) -> int:
        """Extract month number from sheet name (Jan=1, ..., Dec=12)."""
        month_prefix = sheet_name[:3]
        try:
            return month_abbrs.index(month_prefix) + 1
        except ValueError:
            return 999  # Push unrecognized sheets to end

    month_sheets.sort(key=get_month_number)

    logger.info(f"Found {len(month_sheets)} month sheets (sorted by month): {month_sheets}")
    return month_sheets


def find_data_start_row(ws) -> int:
    """Find the first row containing patient data.

    Scans column A for the first numeric value (patient row numbers: 1, 2, 3...).
    This skips any non-numeric values that may appear above the patient data
    (e.g., spaces, text, product data).

    Args:
        ws: openpyxl worksheet object

    Returns:
        Row number (1-indexed) where patient data starts

    Raises:
        ValueError: If no numeric data is found in column A
    """
    max_row = ws.max_row or 1000
    for row_idx in range(1, max_row + 1):
        cell_value = ws.cell(row_idx, 1).value
        if cell_value is not None and isinstance(cell_value, (int, float)):
            return row_idx

    raise ValueError("No patient data found in column A (looking for numeric row numbers)")


def read_header_rows(ws, data_start_row: int, max_cols: int = 100) -> tuple[list, list]:
    """Read and trim the two header rows above the data.

    Headers are located in the two rows immediately before data_start_row.
    Reads up to max_cols columns and trims to the last non-None column.

    Args:
        ws: openpyxl worksheet object
        data_start_row: Row number where patient data starts
        max_cols: Maximum number of columns to read (default: 100)

    Returns:
        Tuple of (header_1, header_2) lists, trimmed to actual width

    Example:
        >>> header_1, header_2 = read_header_rows(ws, 77)
        >>> len(header_1)
        31
    """
    header_row_1 = data_start_row - 1
    header_row_2 = data_start_row - 2

    # Read raw header rows
    header_1_raw = list(
        ws.iter_rows(
            min_row=header_row_1,
            max_row=header_row_1,
            min_col=1,
            max_col=max_cols,
            values_only=True,
        )
    )[0]
    header_2_raw = list(
        ws.iter_rows(
            min_row=header_row_2,
            max_row=header_row_2,
            min_col=1,
            max_col=max_cols,
            values_only=True,
        )
    )[0]

    last_col = max_cols
    for i in range(len(header_1_raw) - 1, -1, -1):
        if header_1_raw[i] is not None or header_2_raw[i] is not None:
            last_col = i + 1
            break

    header_1 = list(header_1_raw[:last_col])
    header_2 = list(header_2_raw[:last_col])

    return header_1, header_2


def merge_headers(header_1: list, header_2: list) -> list[str | None]:
    """Merge two header rows with forward-fill for horizontally merged cells.

    Handles the complex logic of merging multi-line headers while preserving
    information from horizontally merged cells by filling forward.

    Special case: If header_1 contains "Patient ID" (or known synonyms) and
    header_2 appears to be a title row (mostly None), use only header_1.

    Logic:
    - If header_1 contains "Patient ID" and header_2 is mostly None: use header_1 only
    - If both h1 and h2 exist: concatenate as "h2 h1"
    - If only h2 exists: use h2
    - If only h1 exists and both prev_h2 and prev_h1 exist: use "prev_h2 h1" (true horizontal merge)
    - If only h1 exists and prev_h2 but no prev_h1: use h1 (standalone column with header in row 1)
    - If only h1 exists and no prev_h2: use h1
    - If both None: append None

    Args:
        header_1: First header row (closer to data)
        header_2: Second header row (further from data)

    Returns:
        List of merged header strings with whitespace normalized

    Example:
        >>> h1 = ["%", "(dd-mmm-yyyy)", "kg"]
        >>> h2 = ["Updated HbA1c", None, "Body Weight"]
        >>> merge_headers(h1, h2)
        ['Updated HbA1c %', 'Updated HbA1c (dd-mmm-yyyy)', 'Body Weight kg']

        >>> h1 = ["Patient ID", "Patient Name", "Province"]
        >>> h2 = ["Summary of Patient Recruitment", None, None]
        >>> merge_headers(h1, h2)
        ['Patient ID', 'Patient Name', 'Province']
    """
    patient_id_indicators = ["patient id", "patient.id"]
    has_patient_id_in_h1 = any(
        str(h1).strip().lower() in patient_id_indicators
        for h1 in header_1
        if h1 is not None
    )

    non_none_count_h2 = sum(1 for h2 in header_2 if h2 is not None)

    if has_patient_id_in_h1 and non_none_count_h2 <= 2:
        logger.debug(
            "Detected title row in header_2 with Patient ID in header_1, using header_1 only"
        )
        headers = [str(h1).strip() if h1 is not None else None for h1 in header_1]
        headers = [re.sub(r"\s+", " ", h.replace("\n", " ")) if h else None for h in headers]
        return headers

    headers = []
    prev_h2 = None  # Track previous h2 for horizontal merges
    prev_h1 = None  # Track previous h1 to detect true horizontal merges

    for h1, h2 in zip(header_1, header_2, strict=True):
        if h1 and h2:
            headers.append(f"{h2} {h1}".strip())
            prev_h2 = h2
            prev_h1 = h1
        elif h2:
            headers.append(str(h2).strip())
            prev_h2 = h2
            prev_h1 = None
        elif h1:
            # Only forward-fill if previous column also had h1 (true horizontal merge)
            # If prev had h2 but no h1, it's a standalone vertical header
            if prev_h2 and prev_h1:
                headers.append(f"{prev_h2} {h1}".strip())
                prev_h1 = h1
            else:
                headers.append(str(h1).strip())
                prev_h1 = h1
                prev_h2 = None
        else:
            headers.append(None)
            prev_h2 = None
            prev_h1 = None

    headers = [re.sub(r"\s+", " ", h.replace("\n", " ")) if h else None for h in headers]

    return headers


def read_patient_rows(ws, data_start_row: int, num_columns: int) -> list[tuple]:
    """Read patient data rows from the worksheet.

    Reads from data_start_row until either ws.max_row or the first completely
    empty row. Skips rows where both the row number (column A) and patient_id
    (column B) are None, but accepts rows where patient_id exists even if row
    number is missing (handles data quality issues in Excel files).

    Args:
        ws: openpyxl worksheet object
        data_start_row: Row number where patient data starts
        num_columns: Number of columns to read

    Returns:
        List of tuples, each containing one row of patient data

    Example:
        >>> rows = read_patient_rows(ws, 77, 31)
        >>> len(rows)
        4
    """
    data = []
    for row in ws.iter_rows(
        min_row=data_start_row,
        max_row=ws.max_row,
        min_col=1,
        max_col=num_columns,
        values_only=True,
    ):
        if all(cell is None for cell in row):
            break
        # Skip rows where both row number (col A) AND patient_id (col B) are missing
        # This handles cases where Excel has missing row numbers but valid patient data
        if row[0] is None and (len(row) < 2 or row[1] is None):
            continue
        data.append(row)

    return data


def merge_duplicate_columns_data(
    headers: list[str], data: list[list]
) -> tuple[list[str], list[list]]:
    """Merge data from duplicate column headers by concatenating with commas.

    When Excel cells are merged both horizontally and vertically, the forward-fill
    logic in merge_headers() can create duplicate column names. This function
    merges the data from duplicate columns (like R's tidyr::unite()).

    Args:
        headers: List of header strings (may contain duplicates)
        data: List of data rows (each row is a list)

    Returns:
        Tuple of (unique_headers, merged_data)

    Example:
        >>> headers = ["ID", "DM Complications", "DM Complications", "DM Complications", "Age"]
        >>> data = [["1", "A", "B", "C", "25"], ["2", "X", "Y", "Z", "30"]]
        >>> merge_duplicate_columns_data(headers, data)
        (['ID', 'DM Complications', 'Age'], [['1', 'A,B,C', '25'], ['2', 'X,Y,Z', '30']])
    """
    if len(headers) == len(set(headers)):
        return headers, data

    from collections import defaultdict

    header_positions: dict[str, list[int]] = defaultdict(list)
    for idx, header in enumerate(headers):
        header_positions[header].append(idx)

    unique_headers = list(header_positions.keys())

    duplicated = [h for h, positions in header_positions.items() if len(positions) > 1]
    if duplicated:
        logger.debug(f"Merging {len(duplicated)} duplicate column groups: {duplicated}")

    merged_data = []
    for row in data:
        merged_row = []
        for header in unique_headers:
            positions = header_positions[header]
            if len(positions) == 1:
                merged_row.append(row[positions[0]])
            else:
                values = [str(row[pos]) if row[pos] is not None else "" for pos in positions]
                values = [v for v in values if v]
                merged_value = ",".join(values) if values else None
                merged_row.append(merged_value)
        merged_data.append(merged_row)

    return unique_headers, merged_data


def filter_valid_columns(
    headers: list[str | None], data: list[tuple]
) -> tuple[list[str], list[list]]:
    """Filter out columns with None headers and their corresponding data.

    Args:
        headers: List of header strings (may contain None)
        data: List of data rows

    Returns:
        Tuple of (valid_headers, filtered_data)

    Example:
        >>> headers = ["ID", None, "Name", None, "Age"]
        >>> data = [("1", "x", "Alice", "y", "30")]
        >>> filter_valid_columns(headers, data)
        (['ID', 'Name', 'Age'], [['1', 'Alice', '30']])
    """
    valid_cols = [(i, h) for i, h in enumerate(headers) if h]

    if not valid_cols:
        return [], []

    valid_indices = [i for i, _ in valid_cols]
    valid_headers = [h for _, h in valid_cols]

    filtered_data = [[row[i] for i in valid_indices] for row in data]

    return valid_headers, filtered_data


def clean_excel_errors(df: pl.DataFrame) -> pl.DataFrame:
    """Convert Excel error strings to NULL values.

    Excel error codes like #DIV/0!, #VALUE!, etc. are not usable values
    and should be treated as missing data.

    Args:
        df: DataFrame with potential Excel error strings

    Returns:
        DataFrame with Excel errors converted to NULL

    Example:
        >>> df = pl.DataFrame({"bmi": ["17.5", "#DIV/0!", "18.2"]})
        >>> clean_df = clean_excel_errors(df)
        >>> clean_df["bmi"].to_list()
        ['17.5', None, '18.2']
    """
    EXCEL_ERRORS = [
        "#DIV/0!",
        "#VALUE!",
        "#REF!",
        "#NAME?",
        "#NUM!",
        "#N/A",
        "#NULL!",
    ]

    metadata_cols = {"tracker_year", "tracker_month", "clinic_id", "patient_id", "sheet_name", "file_name"}
    data_cols = [col for col in df.columns if col not in metadata_cols]

    if not data_cols:
        return df

    df = df.with_columns([
        pl.when(pl.col(col).is_in(EXCEL_ERRORS))
        .then(None)
        .otherwise(pl.col(col))
        .alias(col)
        for col in data_cols
    ])

    for error in EXCEL_ERRORS:
        for col in data_cols:
            count = (df[col] == error).sum()
            if count > 0:
                logger.debug(f"Converted {count} '{error}' values to NULL in column '{col}'")

    return df


def extract_patient_data(
    tracker_file: Path,
    sheet_name: str,
    year: int,
) -> pl.DataFrame:
    """Extract patient data from a single sheet.

    Orchestrates the extraction process by:
    1. Loading the workbook in read-only mode
    2. Finding where patient data starts
    3. Reading and merging header rows (with forward-fill for horizontal merges)
    4. Filtering valid columns
    5. Reading patient data rows
    6. Creating a Polars DataFrame

    Args:
        tracker_file: Path to the tracker Excel file
        sheet_name: Name of the sheet to extract
        year: Year of the tracker (currently unused, reserved for future use)

    Returns:
        Polars DataFrame with patient data (all columns as strings)

    Example:
        >>> df = extract_patient_data(
        ...     Path("2024_Clinic.xlsx"),
        ...     "Jan24",
        ...     2024
        ... )
        >>> len(df)
        4
        >>> "Patient ID*" in df.columns
        True
    """
    wb = load_workbook(
        tracker_file,
        read_only=True,
        data_only=True,
        keep_vba=False,
        keep_links=False,
    )
    ws = wb[sheet_name]

    data_start_row = find_data_start_row(ws)
    logger.debug(
        f"Sheet '{sheet_name}': Patient data found in rows {data_start_row} to {ws.max_row}"
    )

    logger.info("Processing headers...")
    header_1, header_2 = read_header_rows(ws, data_start_row)
    headers = merge_headers(header_1, header_2)

    valid_cols = [(i, h) for i, h in enumerate(headers) if h]

    if not valid_cols:
        wb.close()
        logger.warning(f"No valid headers found in sheet '{sheet_name}'")
        return pl.DataFrame()

    data = read_patient_rows(ws, data_start_row, len(headers))
    wb.close()

    valid_headers, filtered_data = filter_valid_columns(headers, data)

    valid_headers, filtered_data = merge_duplicate_columns_data(valid_headers, filtered_data)

    # Create DataFrame with ALL columns explicitly as String type to ensure consistent schema
    # across all files and avoid type inference issues (Null vs String dtype)
    df = pl.DataFrame(
        {
            header: pl.Series(
                [str(row[i]) if row[i] is not None else None for row in filtered_data],
                dtype=pl.String,
            )
            for i, header in enumerate(valid_headers)
        }
    )

    logger.info(f"Extracted {len(df)} rows x {len(df.columns)} cols from sheet '{sheet_name}'")

    return df


def harmonize_patient_data_columns(
    df: pl.DataFrame,
    mapper: ColumnMapper | None = None,
    strict: bool = False,
) -> pl.DataFrame:
    """Harmonize patient data columns using synonym mappings.

    Renames columns from their various synonyms (e.g., "Patient ID", "ID",
    "Patient ID*") to standardized column names (e.g., "patient_id").

    Args:
        df: DataFrame with raw column names from tracker
        mapper: ColumnMapper to use (if None, loads default patient mapper)
        strict: If True, raise error if unmapped columns exist
                If False, keep unmapped columns as-is (default)

    Returns:
        DataFrame with standardized column names

    Raises:
        ValueError: If strict=True and unmapped columns exist

    Example:
        >>> raw_df = pl.DataFrame({
        ...     "Patient ID*": ["MY_QI001", "MY_QI002"],
        ...     "Age": [25, 30],
        ... })
        >>> harmonized = harmonize_patient_data_columns(raw_df)
        >>> harmonized.columns
        ['patient_id', 'age']
    """
    if mapper is None:
        mapper = load_patient_mapper()

    renamed_df = mapper.rename_columns(df, strict=strict)

    logger.info(
        f"Harmonized columns: {len(df.columns)} -> {len(renamed_df.columns)} "
        f"({len(df.columns) - len(renamed_df.columns)} columns removed)"
        if len(df.columns) != len(renamed_df.columns)
        else f"Harmonized {len(renamed_df.columns)} columns"
    )

    return renamed_df


def extract_tracker_month(sheet_name: str) -> int:
    """Extract month number (1-12) from sheet name.

    Args:
        sheet_name: Sheet name like "Jan24", "Feb24", etc.

    Returns:
        Month number (1 for January, 2 for February, etc.)

    Raises:
        ValueError: If month cannot be extracted or is out of valid range

    Example:
        >>> extract_tracker_month("Jan24")
        1
        >>> extract_tracker_month("Dec23")
        12
    """
    month_abbrs = list(calendar.month_abbr)[1:]  # ['Jan', 'Feb', ...]

    # Check first 3 characters
    month_prefix = sheet_name[:3]

    if month_prefix in month_abbrs:
        month_num = month_abbrs.index(month_prefix) + 1  # +1 because index is 0-based
        
        # Validate month is in valid range (1-12)
        # This should always be true given the logic above, but check anyway for safety
        if not (1 <= month_num <= 12):
            raise ValueError(
                f"Month number {month_num} is out of valid range (1-12). "
                f"Parsed from sheet name '{sheet_name}'"
            )
        
        return month_num

    raise ValueError(f"Could not extract month from sheet name '{sheet_name}'")


def read_all_patient_sheets(
    tracker_file: Path,
    mapper: ColumnMapper | None = None,
    error_collector: ErrorCollector | None = None,
) -> pl.DataFrame:
    """Read patient data from all month sheets in a tracker file.

    Orchestrates the complete extraction process:
    1. Find all month sheets
    2. Extract tracker year
    3. For each month sheet:
        - Extract raw data
        - Harmonize column names
        - Merge duplicate columns
        - Add metadata (sheet_name, tracker_month, tracker_year, file_name)
    4. Combine all sheets
    5. Filter invalid rows (no patient_id and no name)

    Args:
        tracker_file: Path to the tracker Excel file
        mapper: ColumnMapper to use (if None, loads default patient mapper)
        error_collector: ErrorCollector for tracking data quality issues (optional)

    Returns:
        Combined DataFrame with all patient data from all month sheets

    Raises:
        ValueError: If no month sheets found or year cannot be determined

    Example:
        >>> df = read_all_patient_sheets(Path("2024_Clinic.xlsx"))
        >>> "patient_id" in df.columns
        True
        >>> "tracker_month" in df.columns
        True
        >>> "tracker_year" in df.columns
        True
    """
    logger.info(f"Reading all patient sheets from {tracker_file.name}")

    wb = load_workbook(
        tracker_file, read_only=True, data_only=True, keep_vba=False, keep_links=False
    )

    month_sheets = find_month_sheets(wb)
    if not month_sheets:
        wb.close()
        raise ValueError(f"No month sheets found in {tracker_file.name}")

    year = get_tracker_year(tracker_file, month_sheets)
    logger.info(f"Processing {len(month_sheets)} month sheets for year {year}")

    wb.close()

    all_sheets_data = []

    for sheet_name in month_sheets:
        logger.info(f"Processing sheet: {sheet_name}")

        df_sheet = extract_patient_data(tracker_file, sheet_name, year)

        if df_sheet.is_empty():
            logger.warning(f"Sheet '{sheet_name}' has no data, skipping")
            continue

        df_sheet = harmonize_patient_data_columns(df_sheet, mapper=mapper, strict=False)

        if "patient_id" not in df_sheet.columns:
            logger.warning(
                f"Sheet '{sheet_name}' has no 'patient_id' column after harmonization, skipping"
            )
            continue

        try:
            month_num = extract_tracker_month(sheet_name)
        except ValueError as e:
            logger.warning(f"Could not extract month from '{sheet_name}': {e}, skipping")
            continue

        # Derived metadata (year, month) use Int64; text metadata (sheet_name, etc.) use String
        clinic_id = tracker_file.parent.name
        file_name = tracker_file.stem
        df_sheet = df_sheet.with_columns(
            [
                pl.lit(sheet_name, dtype=pl.String).alias("sheet_name"),
                pl.lit(month_num, dtype=pl.Int64).alias("tracker_month"),
                pl.lit(year, dtype=pl.Int64).alias("tracker_year"),
                pl.lit(file_name, dtype=pl.String).alias("file_name"),
                pl.lit(clinic_id, dtype=pl.String).alias("clinic_id"),
            ]
        )

        all_sheets_data.append(df_sheet)

    if not all_sheets_data:
        raise ValueError(f"No valid patient data found in any month sheets of {tracker_file.name}")

    # Use diagonal_relaxed to handle type mismatches (e.g., Null vs String) like R's bind_rows
    logger.info(f"Combining {len(all_sheets_data)} sheets...")
    df_combined = pl.concat(all_sheets_data, how="diagonal_relaxed")

    initial_rows = len(df_combined)

    # Track rows with missing patient_id for error reporting
    missing_patient_id_rows = df_combined.filter(pl.col("patient_id").is_null())
    missing_count = len(missing_patient_id_rows)

    if missing_count > 0:
        logger.error(
            f"Found {missing_count} rows with missing patient_id in {tracker_file.name} - "
            f"these rows will be excluded from processing"
        )

        # Log to ErrorCollector if available
        if error_collector is not None:
            for row in missing_patient_id_rows.iter_rows(named=True):
                sheet_name = row.get("sheet_name", "unknown")
                name_value = row.get("name", "")
                error_collector.add_error(
                    file_name=tracker_file.stem,
                    patient_id="MISSING",
                    column="patient_id",
                    original_value=None,
                    error_message=f"Row in sheet '{sheet_name}' has missing patient_id (name: {name_value})",
                    error_code="missing_required_field",
                    script="extract",
                    function_name="read_all_patient_sheets",
                )

    # Filter out ALL rows with missing patient_id
    df_combined = df_combined.filter(pl.col("patient_id").is_not_null())

    # Filter out empty rows (both patient_id and name are null/empty) - this is redundant now but kept for clarity
    if "name" in df_combined.columns:
        df_combined = df_combined.filter(
            ~((pl.col("patient_id").str.strip_chars() == "") &
              (pl.col("name").is_null() | (pl.col("name").str.strip_chars() == "")))
        )

    # Filter out rows where both patient_id and name are numeric zeros (0, 0.0, "0", "0.0", etc.)
    if "name" in df_combined.columns:
        df_combined = df_combined.filter(
            ~(pl.col("patient_id").str.strip_chars().is_in(["0", "0.0"]) &
              pl.col("name").str.strip_chars().is_in(["0", "0.0"]))
        )

    # Filter out rows with patient_id starting with "#" (Excel errors like #REF!)
    df_combined = df_combined.filter(~pl.col("patient_id").str.starts_with("#"))

    filtered_rows = initial_rows - len(df_combined)
    if filtered_rows > 0:
        logger.info(f"Filtered out {filtered_rows} invalid rows total")

    df_combined = clean_excel_errors(df_combined)

    wb = load_workbook(
        tracker_file, read_only=True, data_only=True, keep_vba=False, keep_links=False
    )
    all_sheets = wb.sheetnames
    wb.close()

    # Process Patient List sheet if it exists (R: lines 103-130)
    if "Patient List" in all_sheets:
        logger.info("Processing 'Patient List' sheet...")
        try:
            patient_list = extract_patient_data(tracker_file, "Patient List", year)
            if not patient_list.is_empty():
                patient_list = harmonize_patient_data_columns(patient_list, mapper=mapper, strict=False)

                if "patient_id" in patient_list.columns:
                    # Filter out rows with missing patient_id
                    patient_list = patient_list.filter(pl.col("patient_id").is_not_null())

                    # Filter out numeric zeros and Excel errors
                    if "name" in patient_list.columns:
                        patient_list = patient_list.filter(
                            ~(pl.col("patient_id").str.strip_chars().is_in(["0", "0.0"]) &
                              pl.col("name").str.strip_chars().is_in(["0", "0.0"]))
                        )

                    patient_list = patient_list.filter(~pl.col("patient_id").str.starts_with("#"))

                    # R: select(-any_of(c("hba1c_baseline"))) and select(-any_of(c("name")))
                    df_monthly = df_combined.drop("hba1c_baseline") if "hba1c_baseline" in df_combined.columns else df_combined
                    patient_list_join = patient_list.drop("name") if "name" in patient_list.columns else patient_list

                    df_combined = df_monthly.join(
                        patient_list_join,
                        on="patient_id",
                        how="left",
                        suffix=".static"
                    )
                    logger.info(f"Joined {len(patient_list)} Patient List records")
                else:
                    logger.warning("Patient List sheet has no 'patient_id' column after harmonization")
            else:
                logger.warning("Patient List sheet is empty")
        except Exception as e:
            logger.warning(f"Could not process Patient List sheet: {e}")

    # Process Annual sheet if it exists (R: lines 132-160)
    if "Annual" in all_sheets:
        logger.info("Processing 'Annual' sheet...")
        try:
            annual_data = extract_patient_data(tracker_file, "Annual", year)
            if not annual_data.is_empty():
                annual_data = harmonize_patient_data_columns(annual_data, mapper=mapper, strict=False)

                if "patient_id" in annual_data.columns:
                    # Filter out rows with missing patient_id
                    annual_data = annual_data.filter(pl.col("patient_id").is_not_null())

                    # Filter out numeric zeros and Excel errors
                    if "name" in annual_data.columns:
                        annual_data = annual_data.filter(
                            ~(pl.col("patient_id").str.strip_chars().is_in(["0", "0.0"]) &
                              pl.col("name").str.strip_chars().is_in(["0", "0.0"]))
                        )

                    annual_data = annual_data.filter(~pl.col("patient_id").str.starts_with("#"))

                    # R: select(-any_of(c("status", "name")))
                    cols_to_drop = [col for col in ["status", "name"] if col in annual_data.columns]
                    annual_data_join = annual_data.drop(cols_to_drop) if cols_to_drop else annual_data

                    df_combined = df_combined.join(
                        annual_data_join,
                        on="patient_id",
                        how="left",
                        suffix=".annual"
                    )
                    logger.info(f"Joined {len(annual_data)} Annual records")
                else:
                    logger.warning("Annual sheet has no 'patient_id' column after harmonization")
            else:
                logger.warning("Annual sheet is empty")
        except Exception as e:
            logger.warning(f"Could not process Annual sheet: {e}")

    logger.info(
        f"Successfully extracted {len(df_combined)} total rows "
        f"from {len(all_sheets_data)} month sheets"
    )

    # Reorder: metadata first (tracker_year, tracker_month, clinic_id, patient_id), then patient data
    priority_cols = ["tracker_year", "tracker_month", "clinic_id", "patient_id"]
    existing_priority = [c for c in priority_cols if c in df_combined.columns]
    other_cols = [c for c in df_combined.columns if c not in priority_cols]
    df_combined = df_combined.select(existing_priority + other_cols)

    return df_combined


def export_patient_raw(
    df: pl.DataFrame,
    tracker_file: Path,
    output_dir: Path,
) -> Path:
    """Export raw patient data to parquet file.

    Matches R pipeline behavior:
    - Filename: {tracker_name}_patient_raw.parquet
    - Location: output_dir/{tracker_name}_patient_raw.parquet

    Args:
        df: Patient DataFrame to export
        tracker_file: Path to original tracker file (used to extract tracker_name)
        output_dir: Directory to write parquet file (e.g., data_root/output/patient_data_raw)

    Returns:
        Path to the written parquet file

    Example:
        >>> df = read_all_patient_sheets(Path("2024_Clinic.xlsx"))
        >>> output_path = export_patient_raw(
        ...     df,
        ...     Path("2024_Clinic.xlsx"),
        ...     Path("output/patient_data_raw")
        ... )
        >>> output_path.name
        '2024_Clinic_patient_raw.parquet'
    """
    # Extract tracker name (filename without extension)
    tracker_name = tracker_file.stem

    # Create output filename: {tracker_name}_patient_raw.parquet
    output_filename = f"{tracker_name}_patient_raw.parquet"
    output_path = output_dir / output_filename

    # Ensure output directory exists
    output_dir.mkdir(parents=True, exist_ok=True)

    # Write parquet file
    logger.info(f"Writing {len(df)} rows to {output_path}")
    df.write_parquet(output_path)

    logger.info(f"Successfully exported to {output_path}")
    return output_path
